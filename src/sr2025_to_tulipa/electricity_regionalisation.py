from __future__ import annotations

import csv
import math
import re
import unicodedata
from array import array
from dataclasses import asdict, dataclass, replace
from pathlib import Path

from openpyxl import load_workbook

from sr2025_to_tulipa.source_validation import SourceValidationError


SCENARIO_NAMES = {
    "eigen_vermogen": "Eigen Vermogen",
    "gezamenlijke_balans": "Gezamenlijke Balans",
    "horizon_aanvoer": "Horizon Aanvoer",
    "koersvaste_middenweg": "Koersvaste Middenweg",
}
DEMAND_DRIVERS = {
    "agriculture": ("Agriculture", "Agriculture_PtH"),
    "buildings": ("Buildings", "Buildings_hp_electric", "Buildings_hp_hybrid"),
    "bunkers": ("Transport_plane", "Transport_ship"),
    "central_heat": ("District_heat_network", "District_PtH"),
    "households": ("Households", "Households_hp_electric", "Households_hp_hybrid", "Battery_households"),
    "other": ("Other_demand",),
    "transformation": ("CO2_storage", "Datacenters", "Direct_air_capture"),
    "transport": (
        "Transport_bus",
        "Transport_car",
        "Transport_other",
        "Transport_train",
        "Transport_tram",
        "Transport_truck",
        "Transport_van",
        "Battery_transport",
    ),
}
MUNICIPALITY_REPORT_DEMAND_SECTORS = (
    "Agriculture",
    "Buildings",
    "Buildings_hp_electric",
    "Buildings_hp_hybrid",
    "CO2_storage",
    "Datacenters",
    "Direct_air_capture",
    "District_heat_network",
    "Households",
    "Households_hp_electric",
    "Households_hp_hybrid",
    "Other_demand",
    "Transport_bus",
    "Transport_car",
    "Transport_other",
    "Transport_plane",
    "Transport_ship",
    "Transport_train",
    "Transport_tram",
    "Transport_truck",
    "Transport_van",
)
CAPACITY_DRIVERS = {
    "biomass_chp": ("Power_plant_biomass", "Agriculture_CHP"),
    "biomass_chp_dispatchable": ("Power_plant_biomass",),
    "biomass_chp_must_run": ("Power_plant_biomass",),
    "biomass_steam_dispatchable": ("Power_plant_biomass",),
    "biomass_steam_must_run": ("Power_plant_biomass",),
    "coal_steam": ("Power_plant_coal",),
    "coal_steam_ccs": ("Power_plant_coal",),
    "hydro_run_of_river": ("Hydro_RoR",),
    "hydrogen_ccgt": ("Power_plant_hydrogen",),
    "hydrogen_ocgt": ("Power_plant_hydrogen_backup", "Power_plant_hydrogen"),
    "natural_gas_ccgt": ("Power_plant_methane",),
    "natural_gas_chp": ("Power_plant_methane_CHP", "Agriculture_CHP"),
    "natural_gas_ocgt": ("Power_plant_methane_backup", "Power_plant_methane"),
    "natural_gas_steam": ("Power_plant_methane",),
    "nuclear": ("Power_plant_nuclear",),
    "nuclear_smr": ("Power_plant_nuclear_SMR", "Power_plant_nuclear"),
    "solar_pv": ("Solar_PV_buildings", "Solar_PV_field", "Solar_PV_households"),
    "wind_offshore": ("Wind_offshore", "Wind_offshore_hybrid"),
    "wind_onshore": ("Wind_onshore",),
    "waste_chp": ("Power_plant_waste",),
    "waste_chp_ccs": ("Power_plant_waste",),
    "waste_steam": ("Power_plant_waste",),
}
@dataclass(frozen=True)
class Municipality:
    code: str
    name: str
    province: str
    node: str
    industry_power: float


@dataclass(frozen=True)
class SpatialWeightRow:
    scenario_key: str
    year: int
    item: str
    municipality_code: str
    municipality_name: str
    province: str
    node: str
    allocation_share: float
    source_method: str
    reference_year: int


@dataclass(frozen=True)
class NodeCapacityRow:
    scenario_key: str
    year: int
    node: str
    asset_group: str
    fuel: str
    operating_mode: str
    capacity_mw: float
    allocation_share: float
    source_method: str


@dataclass(frozen=True)
class ReconciliationRow:
    scenario_key: str
    year: int
    quantity: str
    item: str
    national_value: float
    distributed_value: float
    difference: float
    status: str


@dataclass(frozen=True)
class DemandSourceReconciliationRow:
    scenario_key: str
    year: int
    reference_year: int
    municipality_demand_mwh: float
    industry_demand_mwh: float
    workbook_total_mwh: float
    mapped_municipality_demand_mwh: float
    unmapped_municipality_demand_mwh: float
    modeled_report_total_mwh: float
    difference_mwh: float
    relative_difference: float
    status: str


@dataclass
class RegionalReference:
    reference_year: int
    municipalities: dict[str, Municipality]
    municipality_volume: dict[str, dict[str, float]]
    municipality_capacity: dict[str, dict[str, float]]
    province_volume: dict[str, dict[str, float]]
    province_capacity: dict[str, dict[str, float]]
    source_municipality_demand_mwh: float | None = None
    source_industry_demand_mwh: float | None = None


def normalize_weights(values: dict[str, float]) -> dict[str, float]:
    """Normalize non-negative allocation drivers and reject an empty driver."""
    cleaned = {key: abs(value) for key, value in values.items() if value != 0.0}
    total = sum(cleaned.values())
    if total <= 0.0:
        raise SourceValidationError("Cannot normalize an empty spatial driver.")
    weights = {key: value / total for key, value in cleaned.items()}
    correction_key = max(weights, key=weights.get)
    weights[correction_key] += 1.0 - sum(weights.values())
    return weights


def aggregate_node_weights(
    municipality_weights: dict[str, float], municipalities: dict[str, Municipality]
) -> dict[str, float]:
    """Aggregate normalized municipality weights to electricity nodes."""
    nodes: dict[str, float] = {}
    for code, weight in municipality_weights.items():
        municipality = municipalities[code]
        nodes[municipality.node] = nodes.get(municipality.node, 0.0) + weight
    return normalize_weights(nodes)


def load_regional_reference(
    scenario_key: str, year: int, regional_dir: Path
) -> RegionalReference:
    """Load one scenario workbook plus the shared municipal and node references."""
    scenario_name = SCENARIO_NAMES[scenario_key]
    reference_year = max(2030, min(2050, year))
    workbook_path = (
        regional_dir
        / f"Scenario {scenario_name}"
        / f"Scenario {scenario_name} {reference_year}.xlsx"
    )
    if not workbook_path.exists():
        raise SourceValidationError(f"Missing regional workbook: {workbook_path}")
    master_path = regional_dir / "Regionalisering vraag SR2025.xlsx"
    municipalities = _load_municipalities(master_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    municipality_volume_sheet = workbook["Municipality (ELEC, volume)"]
    industry_volume_sheet = workbook["Industry (ELEC, volume)"]
    municipality_volume = _read_geography_matrix(municipality_volume_sheet, municipalities)
    municipality_capacity = _read_geography_matrix(
        workbook["Municipality (ELEC, capacity)"], municipalities
    )
    province_volume = _read_geography_matrix(industry_volume_sheet, None)
    province_capacity = _read_geography_matrix(
        workbook["Industry (ELEC, capacity)"], None
    )
    return RegionalReference(
        reference_year=reference_year,
        municipalities=municipalities,
        municipality_volume=municipality_volume,
        municipality_capacity=municipality_capacity,
        province_volume=province_volume,
        province_capacity=province_capacity,
        source_municipality_demand_mwh=_sum_geography_type(
            municipality_volume_sheet, "Demand"
        ),
        source_industry_demand_mwh=_sum_geography_type(
            industry_volume_sheet, "Demand"
        ),
    )


def demand_weights(
    scenario_key: str, year: int, demand_group: str, reference: RegionalReference
) -> list[SpatialWeightRow]:
    """Create municipality demand weights, preserving province-first industry."""
    if demand_group == "industry":
        raw = _province_to_municipal_driver(
            reference.province_volume,
            _industry_demand_sectors(reference),
            reference.municipalities,
        )
        method = "province_industry_then_cbs_municipality_power"
    else:
        sectors = DEMAND_DRIVERS.get(demand_group)
        if sectors is None:
            raw = _generic_municipality_demand(reference)
            method = "municipality_total_demand_proxy"
        else:
            raw = _municipality_driver(reference.municipality_volume, sectors)
            method = "municipality_sector_volume"
            if not any(raw.values()):
                raw = _generic_municipality_demand(reference)
                method = "municipality_total_demand_fallback"
    weights = normalize_weights(raw)
    return [
        SpatialWeightRow(
            scenario_key,
            year,
            demand_group,
            code,
            reference.municipalities[code].name,
            reference.municipalities[code].province,
            reference.municipalities[code].node,
            share,
            method,
            reference.reference_year,
        )
        for code, share in sorted(weights.items())
    ]


def capacity_weights(
    scenario_key: str, year: int, asset_group: str, reference: RegionalReference
) -> list[SpatialWeightRow]:
    """Create technology-specific municipal capacity weights."""
    sectors = CAPACITY_DRIVERS.get(asset_group)
    if sectors is None and asset_group != "hydrogen_chp":
        raise SourceValidationError(
            f"No configured spatial capacity driver for {asset_group}."
        )
    raw = (
        _municipality_driver(reference.municipality_capacity, sectors)
        if sectors
        else {}
    )
    industry_sector = {
        "natural_gas_chp": "Industry_methane_CHP",
        "hydrogen_chp": "Industry_hydrogen_CHP",
    }.get(asset_group)
    if industry_sector:
        industry = _province_to_municipal_driver(
            reference.province_capacity,
            (industry_sector,),
            reference.municipalities,
        )
        raw = {
            code: raw.get(code, 0.0) + industry.get(code, 0.0)
            for code in raw.keys() | industry.keys()
        }
    if sectors and industry_sector:
        method = "municipality_and_province_industry_capacity"
    elif industry_sector:
        method = "province_industry_capacity_then_cbs_municipality_power"
    else:
        method = "municipality_technology_capacity"
    if not any(raw.values()):
        raise SourceValidationError(
            f"No spatial capacity driver for {scenario_key} {year} {asset_group}."
        )
    weights = normalize_weights(raw)
    return [
        SpatialWeightRow(
            scenario_key,
            year,
            asset_group,
            code,
            reference.municipalities[code].name,
            reference.municipalities[code].province,
            reference.municipalities[code].node,
            share,
            method,
            reference.reference_year,
        )
        for code, share in sorted(weights.items())
    ]


def regionalise_demand(
    input_path: Path,
    regional_dir: Path,
    output_dir: Path,
    write_details: bool = True,
) -> None:
    """Distribute national hourly electricity groups through municipalities to nodes."""
    references: dict[tuple[str, int], RegionalReference] = {}
    weight_rows: list[SpatialWeightRow] = []
    node_weights: dict[tuple[str, int, str], dict[str, float]] = {}
    annual_mwh: dict[tuple[str, int, str], float] = {}
    node_profiles: dict[tuple[str, int, str], array] = {}

    with input_path.open(encoding="utf-8", newline="") as input_file:
        for row in csv.DictReader(input_file):
            if row["carrier"] != "electricity":
                continue
            scenario_key = row["scenario_key"]
            year = int(row["year"])
            demand_group = row["demand_group"]
            key = (scenario_key, year, demand_group)
            if key not in node_weights:
                reference_key = (scenario_key, year)
                if reference_key not in references:
                    references[reference_key] = load_regional_reference(
                        scenario_key, year, regional_dir
                    )
                reference = references[reference_key]
                rows = demand_weights(
                    scenario_key, year, demand_group, reference
                )
                weight_rows.extend(rows)
                node_weights[key] = aggregate_node_weights(
                    {item.municipality_code: item.allocation_share for item in rows},
                    reference.municipalities,
                )
            hour = int(row["hour"])
            value = float(row["normalized_mw"])
            annual_mwh[key] = annual_mwh.get(key, 0.0) + value
            for node, share in node_weights[key].items():
                profile = node_profiles.setdefault(
                    (scenario_key, year, node), array("d", [0.0]) * 8760
                )
                profile[hour - 1] += value * share

    for (scenario_key, year), reference in references.items():
        topology_nodes = {item.node for item in reference.municipalities.values()}
        for node in topology_nodes:
            node_profiles.setdefault(
                (scenario_key, year, node), array("d", [0.0]) * 8760
            )

    output_dir.mkdir(parents=True, exist_ok=True)
    if write_details:
        _write_dataclasses(
            weight_rows, output_dir / "electricity_demand_municipality_weights.csv"
        )
    with (output_dir / "electricity_demand_node_hourly.csv").open(
        "w", encoding="utf-8", newline=""
    ) as output_file:
        writer = csv.DictWriter(
            output_file,
            fieldnames=("scenario_key", "year", "node", "hour", "demand_mw"),
        )
        writer.writeheader()
        for (scenario_key, year, node), profile in sorted(node_profiles.items()):
            writer.writerows(
                {
                    "scenario_key": scenario_key,
                    "year": year,
                    "node": node,
                    "hour": hour,
                    "demand_mw": value,
                }
                for hour, value in enumerate(profile, start=1)
            )
    if not write_details:
        return
    municipality_allocations = [
        {
            **asdict(row),
            "annual_mwh": annual_mwh[(row.scenario_key, row.year, row.item)]
            * row.allocation_share,
        }
        for row in weight_rows
    ]
    _write_dicts(
        municipality_allocations,
        output_dir / "electricity_demand_municipality_annual.csv",
    )
    reconciliations = []
    for key, national_value in sorted(annual_mwh.items()):
        distributed_value = national_value * sum(node_weights[key].values())
        reconciliations.append(
            _reconciliation(*key[:2], "demand_mwh", key[2], national_value, distributed_value)
        )
    _write_dataclasses(
        reconciliations, output_dir / "electricity_demand_distribution_reconciliation.csv"
    )
    source_reconciliations = [
        _demand_source_reconciliation(
            scenario_key,
            year,
            reference,
            sum(
                value
                for (row_scenario, row_year, demand_group), value in annual_mwh.items()
                if row_scenario == scenario_key
                and row_year == year
                and demand_group not in {"losses", "power_sector_own_use"}
            ),
        )
        for (scenario_key, year), reference in sorted(references.items())
    ]
    _write_dataclasses(
        source_reconciliations,
        output_dir / "electricity_demand_source_reconciliation.csv",
    )


def regionalise_capacity(
    input_path: Path,
    regional_dir: Path,
    output_dir: Path,
    write_details: bool = True,
) -> None:
    """Distribute grouped national electricity capacity through municipalities to nodes."""
    references: dict[tuple[str, int], RegionalReference] = {}
    all_weights: list[SpatialWeightRow] = []
    output_rows: list[NodeCapacityRow] = []
    reconciliations: list[ReconciliationRow] = []
    with input_path.open(encoding="utf-8", newline="") as input_file:
        for row in csv.DictReader(input_file):
            scenario_key = row["scenario_key"]
            year = int(row["year"])
            asset_group = row["asset_group"]
            item = f"{asset_group}:{row['operating_mode']}"
            reference_key = (scenario_key, year)
            if reference_key not in references:
                references[reference_key] = load_regional_reference(
                    scenario_key, year, regional_dir
                )
            reference = references[reference_key]
            weights = capacity_weights(scenario_key, year, asset_group, reference)
            weights = [replace(weight, item=item) for weight in weights]
            all_weights.extend(weights)
            node_shares = aggregate_node_weights(
                {item.municipality_code: item.allocation_share for item in weights},
                reference.municipalities,
            )
            method = weights[0].source_method
            national_capacity = float(row["capacity_mw"])
            for node, share in sorted(node_shares.items()):
                output_rows.append(
                    NodeCapacityRow(
                        scenario_key,
                        year,
                        node,
                        asset_group,
                        row["fuel"],
                        row["operating_mode"],
                        national_capacity * share,
                        share,
                        method,
                    )
                )
            reconciliations.append(
                _reconciliation(
                    scenario_key,
                    year,
                    "capacity_mw",
                    item,
                    national_capacity,
                    sum(item.capacity_mw for item in output_rows[-len(node_shares) :]),
                )
            )
    output_dir.mkdir(parents=True, exist_ok=True)
    if write_details:
        _write_dataclasses(
            all_weights, output_dir / "electricity_capacity_municipality_weights.csv"
        )
    _write_dataclasses(output_rows, output_dir / "electricity_capacity_by_node.csv")
    if write_details:
        _write_dataclasses(
            reconciliations,
            output_dir / "electricity_capacity_distribution_reconciliation.csv",
        )


def _load_municipalities(master_path: Path) -> dict[str, Municipality]:
    workbook = load_workbook(master_path, read_only=True, data_only=True)
    cbs = workbook["CBS-Dataset-key"]
    municipalities: dict[str, Municipality] = {}
    for values in cbs.iter_rows(min_row=2, values_only=True):
        code = str(values[5] or "")
        node = str(values[10] or "")
        if code.startswith("GM") and node.startswith("E-"):
            municipalities[code] = Municipality(
                code,
                str(values[0]),
                _province_key(str(values[6])),
                node,
                float(values[9] or 0.0),
            )
    mapping_workbook = load_workbook(master_path, read_only=True, data_only=True)
    mapping = mapping_workbook["IO2050 gemeente mapping"]
    by_name = {_name_key(item.name): item for item in municipalities.values()}
    authoritative: dict[str, Municipality] = {}
    for values in mapping.iter_rows(min_row=2, values_only=True):
        code = str(values[1] or "")
        name = str(values[0] or "")
        node = str(values[2] or "")
        if not code.startswith("GM") or not node.startswith("E-"):
            continue
        previous = municipalities.get(code) or by_name.get(_name_key(name))
        authoritative[code] = Municipality(
            code,
            name,
            previous.province if previous else "unknown",
            node,
            previous.industry_power if previous else 0.0,
        )
    overrides_path = master_path.parent / "municipality_node_overrides.csv"
    if overrides_path.exists():
        with overrides_path.open(encoding="utf-8-sig", newline="") as input_file:
            for row in csv.DictReader(input_file):
                code = row["municipality_code"]
                previous = municipalities.get(code)
                authoritative[code] = Municipality(
                    code,
                    row["municipality_name"],
                    _province_key(row["province"]),
                    row["node"],
                    previous.industry_power if previous else 0.0,
                )
    return authoritative


def _read_geography_matrix(worksheet: object, municipalities: dict[str, Municipality] | None) -> dict[str, dict[str, float]]:
    matrix: dict[str, dict[str, float]] = {}
    by_name = (
        {_name_key(item.name): item.code for item in municipalities.values()}
        if municipalities
        else {}
    )
    rows = worksheet.iter_rows(values_only=True)  # type: ignore[attr-defined]
    metadata = [next(rows) for _ in range(6)]
    sectors = metadata[4]
    for values in rows:
        code = str(values[0] or "")
        name = str(values[1] or "")
        if municipalities and code not in municipalities:
            code = by_name.get(_name_key(name), "")
        elif municipalities is None:
            code = _province_key(name)
        if not code:
            continue
        row_values = {
            str(sectors[index]): float(values[index] or 0.0)
            for index in range(2, len(values))
        }
        current = matrix.setdefault(code, {})
        for sector, value in row_values.items():
            current[sector] = current.get(sector, 0.0) + value
    return matrix


def _sum_geography_type(worksheet: object, column_type: str) -> float:
    rows = worksheet.iter_rows(values_only=True)  # type: ignore[attr-defined]
    metadata = [next(rows) for _ in range(6)]
    indexes = [
        index for index, value in enumerate(metadata[3]) if value == column_type
    ]
    return sum(
        abs(float(values[index] or 0.0))
        for values in rows
        for index in indexes
    )


def _municipality_driver(matrix: dict[str, dict[str, float]], sectors: tuple[str, ...]) -> dict[str, float]:
    return {
        code: sum(abs(values.get(sector, 0.0)) for sector in sectors)
        for code, values in matrix.items()
    }


def _industry_demand_sectors(reference: RegionalReference) -> tuple[str, ...]:
    return tuple(
        sorted(
            {
                key
                for values in reference.province_volume.values()
                for key in values
                if key.startswith("Industry_")
                and not any(token in key for token in ("CHP", "DSR", "PtH"))
            }
        )
    )


def _demand_source_reconciliation(
    scenario_key: str,
    year: int,
    reference: RegionalReference,
    modeled_report_total_mwh: float,
) -> DemandSourceReconciliationRow:
    mapped_municipality_total = sum(
        abs(values.get(sector, 0.0))
        for values in reference.municipality_volume.values()
        for sector in MUNICIPALITY_REPORT_DEMAND_SECTORS
    )
    municipality_total = (
        reference.source_municipality_demand_mwh
        if reference.source_municipality_demand_mwh is not None
        else mapped_municipality_total
    )
    industry_total = (
        reference.source_industry_demand_mwh
        if reference.source_industry_demand_mwh is not None
        else sum(
            abs(values.get(sector, 0.0))
            for values in reference.province_volume.values()
            for sector in _industry_demand_sectors(reference)
        )
    )
    workbook_total = municipality_total + industry_total
    difference = modeled_report_total_mwh - workbook_total
    relative_difference = difference / workbook_total if workbook_total else difference
    comparable = reference.reference_year == year
    tolerance = 100_000.0
    status = (
        "pass"
        if comparable and math.isclose(difference, 0.0, abs_tol=tolerance)
        else "reference_year_proxy"
        if not comparable
        else "review"
    )
    return DemandSourceReconciliationRow(
        scenario_key,
        year,
        reference.reference_year,
        municipality_total,
        industry_total,
        workbook_total,
        mapped_municipality_total,
        municipality_total - mapped_municipality_total,
        modeled_report_total_mwh,
        difference,
        relative_difference,
        status,
    )


def _province_to_municipal_driver(
    province_matrix: dict[str, dict[str, float]],
    sectors: tuple[str, ...],
    municipalities: dict[str, Municipality],
) -> dict[str, float]:
    province_totals = {
        _province_key(province): sum(abs(values.get(sector, 0.0)) for sector in set(sectors))
        for province, values in province_matrix.items()
    }
    by_province: dict[str, list[Municipality]] = {}
    for municipality in municipalities.values():
        by_province.setdefault(municipality.province, []).append(municipality)
    result: dict[str, float] = {}
    for province, total in province_totals.items():
        members = by_province.get(province, [])
        if not members or total == 0.0:
            continue
        local_total = sum(abs(item.industry_power) for item in members)
        for item in members:
            local_share = (
                abs(item.industry_power) / local_total
                if local_total
                else 1.0 / len(members)
            )
            result[item.code] = total * local_share
    return result


def _generic_municipality_demand(reference: RegionalReference) -> dict[str, float]:
    sectors = tuple({sector for values in DEMAND_DRIVERS.values() for sector in values})
    raw = _municipality_driver(reference.municipality_volume, sectors)
    if not any(raw.values()):
        return {code: max(item.industry_power, 1.0) for code, item in reference.municipalities.items()}
    return raw


def _province_key(value: str) -> str:
    key = _name_key(value)
    return {"fryslan": "friesland"}.get(key, key)


def _name_key(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", normalized.lower().replace("gemeente", ""))


def _reconciliation(
    scenario_key: str,
    year: int,
    quantity: str,
    item: str,
    national_value: float,
    distributed_value: float,
) -> ReconciliationRow:
    difference = distributed_value - national_value
    tolerance = max(1e-8, abs(national_value) * 1e-10)
    return ReconciliationRow(
        scenario_key,
        year,
        quantity,
        item,
        national_value,
        distributed_value,
        difference,
        "pass" if math.isclose(distributed_value, national_value, abs_tol=tolerance) else "fail",
    )


def _write_dataclasses(rows: list[object], path: Path) -> None:
    _write_dicts([asdict(row) for row in rows], path)  # type: ignore[arg-type]


def _write_dicts(rows: list[dict[str, object]], path: Path) -> None:
    if not rows:
        raise SourceValidationError(f"Cannot write empty regionalisation output {path}.")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as output_file:
        writer = csv.DictWriter(output_file, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
