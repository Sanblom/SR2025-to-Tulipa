from __future__ import annotations

import argparse
import csv
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from sr2025_to_tulipa.electricity_regionalisation import (
    SCENARIO_NAMES,
    _read_geography_matrix,
    load_regional_reference,
    normalize_weights,
)
from sr2025_to_tulipa.etm_client import EtmClient
from sr2025_to_tulipa.source_validation import SourceValidationError


TULIPA_SCHEMA_VERSION = "0.22"
HOURS_PER_YEAR = 8760
TULIPA_SCHEMA_URL = (
    "https://tulipaenergy.github.io/TulipaEnergyModel.jl/v0.22/"
    "20-user-guide/54-input-table-schemas/"
)

ASSET_COLUMNS = (
    "asset", "capacity", "capacity_storage_energy", "consumer_balance_sense",
    "discount_rate", "economic_lifetime", "energy_to_power_ratio",
    "investment_integer", "investment_integer_storage_energy", "is_seasonal",
    "max_ramp_down", "max_ramp_up", "min_operating_point", "ramping",
    "storage_method_energy", "technical_lifetime", "type", "unit_commitment",
    "unit_commitment_integer", "use_binary_storage_method", "vintage_method",
)
ASSET_BOTH_COLUMNS = (
    "asset", "commission_year", "decommissionable", "initial_storage_units",
    "initial_units", "milestone_year",
)
ASSET_COMMISSION_COLUMNS = (
    "asset", "commission_year", "conversion_efficiency", "fixed_cost",
    "fixed_cost_storage_energy", "investment_cost", "investment_cost_storage_energy",
    "investment_limit", "investment_limit_storage_energy",
    "storage_charging_efficiency", "storage_discharging_efficiency",
    "storage_loss_from_stored_energy",
)
ASSET_MILESTONE_COLUMNS = (
    "asset", "initial_storage_level", "investable", "max_energy_timeframe_partition",
    "milestone_year", "min_energy_timeframe_partition", "peak_demand",
    "storage_inflows", "units_on_cost",
)
FLOW_COLUMNS = (
    "capacity", "carrier", "discount_rate", "economic_lifetime", "from_asset",
    "investment_integer", "is_transport", "technical_lifetime", "to_asset",
)
FLOW_BOTH_COLUMNS = (
    "commission_year", "decommissionable", "from_asset", "initial_export_units",
    "initial_import_units", "milestone_year", "to_asset",
)
FLOW_COMMISSION_COLUMNS = (
    "capacity_coefficient", "commission_year", "conversion_coefficient", "fixed_cost",
    "from_asset", "investment_cost", "investment_limit", "producer_efficiency",
    "to_asset",
)
FLOW_MILESTONE_COLUMNS = (
    "commodity_price", "dc_opf", "from_asset", "investable", "milestone_year",
    "operational_cost", "reactance", "to_asset",
)

TABLE_COLUMNS = {
    "asset": ASSET_COLUMNS,
    "asset-both": ASSET_BOTH_COLUMNS,
    "asset-commission": ASSET_COMMISSION_COLUMNS,
    "asset-milestone": ASSET_MILESTONE_COLUMNS,
    "flow": FLOW_COLUMNS,
    "flow-both": FLOW_BOTH_COLUMNS,
    "flow-commission": FLOW_COMMISSION_COLUMNS,
    "flow-milestone": FLOW_MILESTONE_COLUMNS,
}

HYDROGEN_PLACEMENT_DRIVERS = {
    "electrolysis": (
        "Power_to_gas_onshore",
        "Power_to_gas_wind_offshore_hybrid",
    ),
    "atr_ccs": ("Blue_hydrogen_production",),
    "smr": ("Grey_hydrogen_production",),
    "smr_ccs": ("Grey_hydrogen_production_CCS",),
}

ELECTRICITY_STORAGE_PLACEMENT_DRIVERS = {
    "battery_system": ("Battery_system", "Battery_solar_PV", "Battery_wind_onshore"),
    "ides_storage": ("IDES_storage",),
    "mdes_storage": ("MDES_storage",),
}

ELECTRICITY_STORAGE_QUERIES = {
    "battery_system": "energy_flexibility_mv_batteries",
    "ides_storage": "energy_flexibility_flow_batteries",
    "mdes_storage": "energy_flexibility_hv_opac",
}

HYDROGEN_CARRIER_IMPORT_COST_INPUTS = {
    "ammonia_reformer": "costs_imported_ammonia",
    "liquid_hydrogen_regasifier": "costs_imported_liquid_hydrogen",
    "lohc_reformer": "costs_imported_lohc",
}

HYDROGEN_CARRIER_IMPORT_EFFICIENCIES = {
    "liquid_hydrogen_regasifier": 0.986,
    "lohc_reformer": 0.694,
}

DIRECT_HYDROGEN_IMPORT_GROUPS = {
    "hydrogen_import_backup",
    "hydrogen_import_baseload",
}

GENERATION_FUELS = {
    "biomass", "coal", "diesel", "hydro", "nuclear", "solar", "waste", "wind",
}

POWER_CONVERSION_EFFICIENCY_INPUTS = {
    "hydrogen_ccgt": "efficiency_energy_power_combined_cycle_hydrogen",
    "hydrogen_chp": "external_coupling_efficiency_industry_chp_turbine_hydrogen_electricity",
    "hydrogen_ocgt": "efficiency_energy_power_turbine_hydrogen",
    "natural_gas_ccgt": "efficiency_energy_power_combined_cycle_network_gas",
    "natural_gas_ocgt": "efficiency_energy_power_turbine_network_gas",
    "natural_gas_steam": "efficiency_energy_power_ultra_supercritical_network_gas",
}

NATURAL_GAS_CHP_EFFICIENCY_INPUTS = {
    "merit_order_agriculture_chp_gas_dispatchable_capacity_in_merit_order_table":
        "efficiency_agriculture_chp_engine_dispatchable_network_gas_electricity",
    "merit_order_gas_chp_engine_mt_capacity_in_merit_order_table":
        "efficiency_energy_chp_local_engine_network_gas_electricity",
    "merit_order_gas_chp_ccgt_mt_capacity_in_merit_order_table":
        "efficiency_energy_chp_combined_cycle_network_gas_electricity",
    "merit_order_industry_chp_gas_turbine_capacity_in_merit_order_table":
        "external_coupling_efficiency_industry_chp_turbine_gas_power_fuelmix_electricity",
}

HYDROGEN_CONVERSION_ENERGY_QUERIES = {
    "atr_ccs": (
        "natural_gas_for_autothermal_reformer_ccs_conversion",
        "autothermal_reformer_ccs_in_source_of_hydrogen_production",
    ),
    "smr": (
        "natural_gas_for_steam_methane_reformer_conversion",
        "steam_methane_reformer_in_source_of_hydrogen_production",
    ),
    "smr_ccs": (
        "natural_gas_for_steam_methane_reformer_ccs_conversion",
        "steam_methane_reformer_ccs_in_source_of_hydrogen_production",
    ),
}

I_ELGAS_REFORMER_MAPPING = {
    "atr_ccs": "ATR",
    "smr": "SMR",
    "smr_ccs": "SMR_CCS",
}

TYNDP_PROFILE_SOURCES = {
    "wind_onshore": ("NL00_Wind_Onshore",),
    "wind_offshore": ("NL00_Wind_Offshore",),
    "solar_pv": ("NL00_Solar_Photovoltaic", "NL00_Solar_Rooftop"),
    "hydro_run_of_river": ("NL00_Hydro_Run_of_River",),
}

TYNDP_PROFILE_DIRECTORIES = {
    2025: "tulipa_input_north_sea_2026_2030",
    2030: "tulipa_input_north_sea_2026_2030",
    2035: "tulipa_input_north_sea_2026_2035",
    2040: "tulipa_input_north_sea_2026",
    2050: "tulipa_input_north_sea_2026_2050",
}

MUST_RUN_CURVE_SOURCES = {
    "biomass_chp_must_run": (
        "energy_chp_local_engine_mt_biogas",
        "energy_chp_local_mt_wood_pellets_must_run",
    ),
    "biomass_steam_must_run": ("energy_power_wood_pellets_must_run",),
    "waste_chp": ("energy_chp_supercritical_ht_waste_mix",),
    "waste_chp_ccs": ("energy_chp_supercritical_ccs_ht_waste_mix",),
    "waste_steam": ("energy_power_supercritical_waste_mix",),
}


@dataclass(frozen=True)
class ConversionEfficiencyRow:
    scenario_key: str
    year: int
    asset_group: str
    input_capacity_mw: float
    output_capacity_mw: float
    conversion_efficiency: float
    capacity_basis: str


@dataclass(frozen=True)
class TechnologyParameter:
    asset_group: str
    conversion_efficiency: float
    variable_cost_eur_per_mwh_output: float
    efficiency_source: str
    cost_source: str
    non_fuel_variable_cost_eur_per_mwh_output: float = 0.0
    emissions_kg_per_gj_input: float = 0.0
    co2_cost_eur_per_mwh_output: float = 0.0
    co2_price_eur_per_tonne: float = 0.0


@dataclass(frozen=True)
class IElgasTechnologyParameter:
    technology: str
    emissions_kg_per_gj_input: float
    marginal_cost_eur_per_mwh_output: float


def reformer_co2_cost_eur_per_mwh_output(
    emissions_kg_per_gj_input: float,
    conversion_efficiency: float,
    co2_price_eur_per_tonne: float,
) -> float:
    if not 0.0 < conversion_efficiency <= 1.0:
        raise SourceValidationError(
            f"Invalid reformer conversion efficiency: {conversion_efficiency}."
        )
    return (
        emissions_kg_per_gj_input
        * 3.6
        / 1000.0
        / conversion_efficiency
        * co2_price_eur_per_tonne
    )


@dataclass(frozen=True)
class StorageTechnologyRow:
    asset_group: str
    carrier: str
    input_capacity_mw: float
    output_capacity_mw: float
    storage_volume_mwh: float
    charging_efficiency: float
    discharging_efficiency: float
    is_seasonal: bool


def electricity_bus(node: str, year: int) -> str:
    return f"{node}_E_Demand_{year}"


def electricity_endpoint_asset(endpoint: str, year: int) -> str:
    return electricity_bus(endpoint, year)


def country_bus(carrier: str, year: int) -> str:
    abbreviation = {"hydrogen": "H", "methane": "M"}[carrier]
    return f"NL_{abbreviation}_Demand_{year}"


def _asset_row(
    asset: str,
    asset_type: str,
    capacity: float = 0.0,
    storage_volume_mwh: float = 0.0,
    is_seasonal: bool = False,
) -> dict[str, object]:
    energy_to_power_ratio = storage_volume_mwh / capacity if capacity > 0.0 else 0.0
    return {
        "asset": asset,
        "capacity": capacity,
        "capacity_storage_energy": storage_volume_mwh,
        "consumer_balance_sense": "==",
        "discount_rate": 0.0,
        "economic_lifetime": 1,
        "energy_to_power_ratio": energy_to_power_ratio,
        "investment_integer": False,
        "investment_integer_storage_energy": False,
        "is_seasonal": is_seasonal,
        "max_ramp_down": 0.0,
        "max_ramp_up": 0.0,
        "min_operating_point": 0.0,
        "ramping": False,
        "storage_method_energy": (
            "use_fixed_energy_to_power_ratio" if asset_type == "storage" else "none"
        ),
        "technical_lifetime": 1,
        "type": asset_type,
        "unit_commitment": "none",
        "unit_commitment_integer": False,
        "use_binary_storage_method": "",
        "vintage_method": "aggregated",
    }


def _add_asset(
    frames: dict[str, list[dict[str, object]]],
    asset: str,
    asset_type: str,
    capacity: float,
    year: int,
    conversion_efficiency: float = 1.0,
    storage_volume_mwh: float = 0.0,
    charging_efficiency: float = 1.0,
    discharging_efficiency: float = 1.0,
    is_seasonal: bool = False,
) -> None:
    frames["asset"].append(
        _asset_row(asset, asset_type, capacity, storage_volume_mwh, is_seasonal)
    )
    frames["asset-both"].append({
        "asset": asset,
        "commission_year": year,
        "decommissionable": False,
        "initial_storage_units": 0.0,
        "initial_units": 1.0,
        "milestone_year": year,
    })
    frames["asset-commission"].append({
        "asset": asset,
        "commission_year": year,
        "conversion_efficiency": conversion_efficiency,
        "fixed_cost": 0.0,
        "fixed_cost_storage_energy": 0.0,
        "investment_cost": 0.0,
        "investment_cost_storage_energy": 0.0,
        "investment_limit": "",
        "investment_limit_storage_energy": "",
        "storage_charging_efficiency": charging_efficiency,
        "storage_discharging_efficiency": discharging_efficiency,
        "storage_loss_from_stored_energy": 0.0,
    })
    frames["asset-milestone"].append({
        "asset": asset,
        "initial_storage_level": "",
        "investable": False,
        "max_energy_timeframe_partition": "",
        "milestone_year": year,
        "min_energy_timeframe_partition": "",
        "peak_demand": 0.0,
        "storage_inflows": 0.0,
        "units_on_cost": "",
    })


def _add_flow(
    frames: dict[str, list[dict[str, object]]],
    from_asset: str,
    to_asset: str,
    carrier: str,
    year: int,
    capacity_coefficient: float = 1.0,
    operational_cost: float = 0.0,
    commodity_price: float = 0.0,
) -> None:
    frames["flow"].append({
        "capacity": "",
        "carrier": carrier,
        "discount_rate": 0.0,
        "economic_lifetime": 1,
        "from_asset": from_asset,
        "investment_integer": False,
        "is_transport": False,
        "technical_lifetime": 1,
        "to_asset": to_asset,
    })
    frames["flow-commission"].append({
        "capacity_coefficient": capacity_coefficient,
        "commission_year": year,
        "conversion_coefficient": 1.0,
        "fixed_cost": 0.0,
        "from_asset": from_asset,
        "investment_cost": 0.0,
        "investment_limit": "",
        "producer_efficiency": 1.0,
        "to_asset": to_asset,
    })
    frames["flow-milestone"].append({
        "commodity_price": commodity_price,
        "dc_opf": False,
        "from_asset": from_asset,
        "investable": False,
        "milestone_year": year,
        "operational_cost": operational_cost,
        "reactance": "",
        "to_asset": to_asset,
    })


def _add_transport_flow(
    frames: dict[str, list[dict[str, object]]],
    from_asset: str,
    to_asset: str,
    forward_capacity_mw: float,
    reverse_capacity_mw: float,
    year: int,
) -> None:
    base_capacity = (
        forward_capacity_mw if forward_capacity_mw > 0.0 else reverse_capacity_mw
    )
    if base_capacity <= 0.0:
        return
    frames["flow"].append({
        "capacity": base_capacity,
        "carrier": "electricity",
        "discount_rate": 0.0,
        "economic_lifetime": 1,
        "from_asset": from_asset,
        "investment_integer": False,
        "is_transport": True,
        "technical_lifetime": 1,
        "to_asset": to_asset,
    })
    frames["flow-both"].append({
        "commission_year": year,
        "decommissionable": False,
        "from_asset": from_asset,
        "initial_export_units": forward_capacity_mw / base_capacity,
        "initial_import_units": reverse_capacity_mw / base_capacity,
        "milestone_year": year,
        "to_asset": to_asset,
    })
    frames["flow-commission"].append({
        "capacity_coefficient": 1.0,
        "commission_year": year,
        "conversion_coefficient": 1.0,
        "fixed_cost": 0.0,
        "from_asset": from_asset,
        "investment_cost": 0.0,
        "investment_limit": "",
        "producer_efficiency": 1.0,
        "to_asset": to_asset,
    })
    frames["flow-milestone"].append({
        "commodity_price": 0.0,
        "dc_opf": False,
        "from_asset": from_asset,
        "investable": False,
        "milestone_year": year,
        "operational_cost": 0.0,
        "reactance": 0.3,
        "to_asset": to_asset,
    })


def assemble_tulipa_tables(
    year: int,
    electricity_rows: Iterable[dict[str, object]],
    hydrogen_rows: Iterable[dict[str, object]],
    hydrogen_node_weights: dict[str, dict[str, float]],
    electricity_nodes: Iterable[str] | None = None,
    transport_links: Iterable[dict[str, object]] = (),
    technology_parameters: dict[str, TechnologyParameter] | None = None,
    storage_rows: Iterable[StorageTechnologyRow] = (),
    electricity_storage_weights: dict[str, dict[str, float]] | None = None,
    generation_costs_eur_per_mwh: dict[str, float] | None = None,
    hydrogen_carrier_import_costs_eur_per_mwh: dict[str, float] | None = None,
    hydrogen_import_costs_eur_per_mwh: dict[str, float] | None = None,
    hydrogen_carrier_production_twh: dict[tuple[str, str], float] | None = None,
    methane_rows: Iterable[dict[str, object]] = (),
    fuel_prices_eur_per_mwh: dict[str, float] | None = None,
) -> tuple[
    dict[str, list[dict[str, object]]],
    list[dict[str, object]],
    list[dict[str, object]],
]:
    electricity_rows = list(electricity_rows)
    hydrogen_rows = list(hydrogen_rows)
    technology_parameters = technology_parameters or {}
    storage_rows = list(storage_rows)
    electricity_storage_weights = electricity_storage_weights or {}
    generation_costs_eur_per_mwh = generation_costs_eur_per_mwh or {}
    hydrogen_carrier_import_costs_eur_per_mwh = (
        hydrogen_carrier_import_costs_eur_per_mwh or {}
    )
    hydrogen_import_costs_eur_per_mwh = hydrogen_import_costs_eur_per_mwh or {}
    hydrogen_carrier_production_twh = hydrogen_carrier_production_twh or {}
    methane_rows = list(methane_rows)
    fuel_prices_eur_per_mwh = fuel_prices_eur_per_mwh or {}
    nodes = sorted(
        set(electricity_nodes or ())
        or {str(row["node"]) for row in electricity_rows}
    )
    if not nodes:
        raise SourceValidationError("No electricity nodes found for Tulipa asset generation.")

    frames = {table: [] for table in TABLE_COLUMNS}
    reconciliations: list[dict[str, object]] = []
    storage_reconciliations: list[dict[str, object]] = []
    for node in nodes:
        _add_asset(frames, electricity_bus(node, year), "consumer", 0.0, year)
    for carrier in ("hydrogen", "methane"):
        _add_asset(frames, country_bus(carrier, year), "consumer", 0.0, year)

    for row in methane_rows:
        group = str(row["supply_group"])
        if group == "network_gas_storage":
            continue
        capacity = float(row["summed_route_peak_mw"])
        annual_energy = float(row["annual_supply_twh"]) * 1_000_000.0
        if capacity <= 0.0 or annual_energy <= 0.0:
            continue
        asset = f"NL_{group}_{year}"
        _add_asset(frames, asset, "producer", capacity, year)
        frames["asset-milestone"][-1]["max_energy_timeframe_partition"] = annual_energy
        commodity_price = (
            fuel_prices_eur_per_mwh["renewable_methane"]
            if group == "renewable_methane"
            else fuel_prices_eur_per_mwh["methane"]
        )
        _add_flow(
            frames,
            asset,
            country_bus("methane", year),
            "methane",
            year,
            commodity_price=commodity_price,
        )

    external_countries = sorted({
        str(link[endpoint])
        for link in transport_links
        for endpoint in ("from_endpoint", "to_endpoint")
        if not str(link[endpoint]).startswith("E-")
    })
    for country in external_countries:
        _add_asset(
            frames, electricity_endpoint_asset(country, year), "consumer", 0.0, year
        )

    for link in transport_links:
        _add_transport_flow(
            frames,
            electricity_endpoint_asset(str(link["from_endpoint"]), year),
            electricity_endpoint_asset(str(link["to_endpoint"]), year),
            float(link["forward_capacity_mw"]),
            float(link["reverse_capacity_mw"]),
            year,
        )

    for row in electricity_rows:
        capacity = float(row["capacity_mw"])
        fuel = str(row["fuel"])
        if capacity <= 0.0:
            continue
        node = str(row["node"])
        group = str(row["asset_group"])
        mode = str(row["operating_mode"])
        asset = f"{node}_{group}_{mode}"
        if fuel in {"hydrogen", "natural_gas"}:
            parameters = technology_parameters[group]
            _add_asset(
                frames,
                asset,
                "conversion",
                capacity,
                year,
                parameters.conversion_efficiency,
            )
            input_carrier = "hydrogen" if fuel == "hydrogen" else "methane"
            _add_flow(
                frames, country_bus(input_carrier, year), asset, input_carrier, year
            )
        else:
            _add_asset(frames, asset, "producer", capacity, year)
        _add_flow(
            frames,
            asset,
            electricity_bus(node, year),
            "electricity",
            year,
            operational_cost=(
                technology_parameters[group].variable_cost_eur_per_mwh_output
                if fuel in {"hydrogen", "natural_gas"}
                else generation_costs_eur_per_mwh.get(group, 0.0)
            ),
        )

    for row in hydrogen_rows:
        group = str(row["asset_group"])
        if group not in HYDROGEN_PLACEMENT_DRIVERS:
            continue
        national_capacity = float(row["capacity_mw"])
        weights = hydrogen_node_weights.get(group, {})
        if national_capacity > 0.0 and not weights:
            raise SourceValidationError(f"No E-node placement weights for {group}.")
        distributed = 0.0
        for node, share in sorted(weights.items()):
            capacity = national_capacity * share
            if capacity <= 0.0:
                continue
            mode = str(row["operating_mode"])
            asset = f"{node}_{group}_{mode}"
            parameters = technology_parameters[group]
            efficiency = parameters.conversion_efficiency
            _add_asset(
                frames, asset, "conversion", capacity, year, efficiency
            )
            input_carrier = "electricity" if group == "electrolysis" else "methane"
            input_bus = (
                electricity_bus(node, year)
                if input_carrier == "electricity"
                else country_bus("methane", year)
            )
            _add_flow(
                frames,
                input_bus,
                asset,
                input_carrier,
                year,
            )
            _add_flow(
                frames,
                asset,
                country_bus("hydrogen", year),
                "hydrogen",
                year,
                operational_cost=parameters.variable_cost_eur_per_mwh_output,
            )
            distributed += capacity
        reconciliations.append({
            "asset_group": group,
            "operating_mode": row["operating_mode"],
            "national_capacity_mw": national_capacity,
            "distributed_capacity_mw": distributed,
            "difference_mw": distributed - national_capacity,
        })

    for row in hydrogen_rows:
        group = str(row["asset_group"])
        if group not in HYDROGEN_CARRIER_IMPORT_COST_INPUTS:
            continue
        capacity = float(row["capacity_mw"])
        if capacity <= 0.0:
            continue
        mode = str(row["operating_mode"])
        asset = f"NL_{group}_{mode}_{year}"
        _add_asset(frames, asset, "producer", capacity, year)
        if mode == "must_run":
            frames["asset-milestone"][-1]["min_energy_timeframe_partition"] = (
                hydrogen_carrier_production_twh[(group, mode)] * 1_000_000.0
            )
        _add_flow(
            frames,
            asset,
            country_bus("hydrogen", year),
            "hydrogen",
            year,
            operational_cost=hydrogen_carrier_import_costs_eur_per_mwh[group],
        )
        reconciliations.append({
            "asset_group": group,
            "operating_mode": mode,
            "national_capacity_mw": capacity,
            "distributed_capacity_mw": capacity,
            "difference_mw": 0.0,
        })

    for row in hydrogen_rows:
        group = str(row["asset_group"])
        if group not in DIRECT_HYDROGEN_IMPORT_GROUPS:
            continue
        capacity = float(row["capacity_mw"])
        if capacity <= 0.0:
            continue
        asset = f"NL_{group}_{year}"
        _add_asset(frames, asset, "producer", capacity, year)
        _add_flow(
            frames,
            asset,
            country_bus("hydrogen", year),
            "hydrogen",
            year,
            operational_cost=hydrogen_import_costs_eur_per_mwh[group],
        )
        if group == "hydrogen_import_baseload":
            frames["asset-milestone"][-1][
                "min_energy_timeframe_partition"
            ] = capacity * 8760.0
        reconciliations.append({
            "asset_group": group,
            "operating_mode": row["operating_mode"],
            "national_capacity_mw": capacity,
            "distributed_capacity_mw": capacity,
            "difference_mw": 0.0,
        })

    for row in storage_rows:
        if row.output_capacity_mw <= 0.0 or row.storage_volume_mwh <= 0.0:
            continue
        if row.carrier == "electricity":
            weights = electricity_storage_weights.get(row.asset_group, {})
            if not weights:
                raise SourceValidationError(
                    f"No E-node storage placement weights for {row.asset_group}."
                )
            placements = [
                (node, electricity_bus(node, year), share)
                for node, share in sorted(weights.items())
            ]
        else:
            placements = [("NL", country_bus("hydrogen", year), 1.0)]

        distributed_input = 0.0
        distributed_output = 0.0
        distributed_volume = 0.0
        for node, bus, share in placements:
            output_capacity = row.output_capacity_mw * share
            input_capacity = row.input_capacity_mw * share
            storage_volume = row.storage_volume_mwh * share
            if output_capacity <= 0.0:
                continue
            asset = f"{node}_{row.asset_group}_{year}"
            _add_asset(
                frames,
                asset,
                "storage",
                output_capacity,
                year,
                storage_volume_mwh=storage_volume,
                charging_efficiency=row.charging_efficiency,
                discharging_efficiency=row.discharging_efficiency,
                is_seasonal=row.is_seasonal,
            )
            _add_flow(
                frames,
                bus,
                asset,
                row.carrier,
                year,
                row.output_capacity_mw / row.input_capacity_mw,
            )
            _add_flow(frames, asset, bus, row.carrier, year)
            distributed_input += input_capacity
            distributed_output += output_capacity
            distributed_volume += storage_volume
        storage_reconciliations.append({
            "asset_group": row.asset_group,
            "carrier": row.carrier,
            "source_input_capacity_mw": row.input_capacity_mw,
            "tulipa_input_capacity_mw": distributed_input,
            "input_difference_mw": distributed_input - row.input_capacity_mw,
            "source_output_capacity_mw": row.output_capacity_mw,
            "tulipa_output_capacity_mw": distributed_output,
            "output_difference_mw": distributed_output - row.output_capacity_mw,
            "source_storage_volume_mwh": row.storage_volume_mwh,
            "tulipa_storage_volume_mwh": distributed_volume,
            "storage_volume_difference_mwh": distributed_volume - row.storage_volume_mwh,
        })

    return frames, reconciliations, storage_reconciliations


def load_case_rows(path: Path, scenario_key: str, year: int) -> list[dict[str, object]]:
    with path.open(encoding="utf-8", newline="") as input_file:
        return [
            dict(row)
            for row in csv.DictReader(input_file)
            if row["scenario_key"] == scenario_key and int(row["year"]) == year
        ]


def load_electricity_transport_links(
    workbook_path: Path, model_year: int
) -> list[dict[str, object]]:
    source_year = 2027 if model_year == 2025 else model_year
    if source_year not in {2027, 2030, 2035, 2040, 2050}:
        raise SourceValidationError(
            f"No I-ELGAS electricity trading capacities for {model_year}."
        )
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook["Inputs I-ELGAS DE"]
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    normalized_headers = tuple(str(value).strip() for value in headers)
    year_index = normalized_headers.index(str(source_year))
    directional: dict[tuple[str, str], float] = {}
    for values in rows:
        first = str(values[0] or "").strip()
        second = str(values[1] or "").strip()
        first_is_electricity = first.startswith("E-")
        second_is_electricity = second.startswith("E-")
        if not (first_is_electricity or second_is_electricity):
            continue
        if first.startswith(("H-", "G-")) or second.startswith(("H-", "G-")):
            continue
        capacity = float(values[year_index] or 0.0)
        directional[(first, second)] = directional.get((first, second), 0.0) + capacity

    links: list[dict[str, object]] = []
    for first, second in sorted({tuple(sorted(pair)) for pair in directional}):
        forward = directional.get((first, second), 0.0)
        reverse = directional.get((second, first), 0.0)
        if forward <= 0.0 and reverse <= 0.0:
            continue
        links.append({
            "from_endpoint": first,
            "to_endpoint": second,
            "forward_capacity_mw": forward,
            "reverse_capacity_mw": reverse,
            "source_year": source_year,
            "link_scope": (
                "internal_nl"
                if first.startswith("E-") and second.startswith("E-")
                else "nl_cross_border"
            ),
        })
    return links


def load_hydrogen_node_weights(
    scenario_key: str, year: int, regional_dir: Path
) -> tuple[dict[str, dict[str, float]], set[str]]:
    reference = load_regional_reference(scenario_key, year, regional_dir)
    scenario_name = SCENARIO_NAMES[scenario_key]
    reference_year = max(2030, min(2050, year))
    workbook_path = (
        regional_dir
        / f"Scenario {scenario_name}"
        / f"Scenario {scenario_name} {reference_year}.xlsx"
    )
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    matrix = _read_geography_matrix(
        workbook["Municipality (H2, capacity)"], reference.municipalities
    )
    result: dict[str, dict[str, float]] = {}
    for group, drivers in HYDROGEN_PLACEMENT_DRIVERS.items():
        raw: dict[str, float] = {}
        for code, values in matrix.items():
            node = reference.municipalities[code].node
            raw[node] = raw.get(node, 0.0) + sum(abs(values.get(driver, 0.0)) for driver in drivers)
        positive = {node: value for node, value in raw.items() if value > 0.0}
        if positive:
            result[group] = normalize_weights(positive)
    electricity_nodes = {item.node for item in reference.municipalities.values()}
    return result, electricity_nodes


def load_electricity_storage_weights(
    scenario_key: str, year: int, regional_dir: Path
) -> dict[str, dict[str, float]]:
    reference = load_regional_reference(scenario_key, year, regional_dir)
    scenario_name = SCENARIO_NAMES[scenario_key]
    reference_year = max(2030, min(2050, year))
    workbook_path = (
        regional_dir
        / f"Scenario {scenario_name}"
        / f"Scenario {scenario_name} {reference_year}.xlsx"
    )
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    matrix = _read_geography_matrix(
        workbook["Municipality (ELEC, capacity)"], reference.municipalities
    )
    result: dict[str, dict[str, float]] = {}
    for group, drivers in ELECTRICITY_STORAGE_PLACEMENT_DRIVERS.items():
        raw: dict[str, float] = {}
        for code, values in matrix.items():
            node = reference.municipalities[code].node
            raw[node] = raw.get(node, 0.0) + sum(
                abs(values.get(driver, 0.0)) for driver in drivers
            )
        positive = {node: value for node, value in raw.items() if value > 0.0}
        if positive:
            result[group] = normalize_weights(positive)
    return result


def _future_value(values: dict[str, object], query_key: str) -> float:
    result = values[query_key]
    if not isinstance(result, dict) or result.get("future") is None:
        raise SourceValidationError(f"ETM storage query {query_key} has no future value.")
    return float(result["future"])


def aggregate_generation_costs(
    mappings: Iterable[dict[str, str]], values: dict[str, object]
) -> dict[str, float]:
    weighted_costs: dict[str, float] = {}
    capacities: dict[str, float] = {}
    for mapping in mappings:
        if mapping["included"].lower() != "true":
            continue
        capacity_query = mapping["query_key"]
        cost_query = capacity_query.replace(
            "_capacity_in_merit_order_table",
            "_operating_costs_in_merit_order_table",
        )
        capacity = _future_value(values, capacity_query)
        if capacity <= 0.0:
            continue
        group = mapping["asset_group"]
        weighted_costs[group] = (
            weighted_costs.get(group, 0.0)
            + capacity * _future_value(values, cost_query)
        )
        capacities[group] = capacities.get(group, 0.0) + capacity
    return {
        group: weighted_cost / capacities[group]
        for group, weighted_cost in weighted_costs.items()
    }


def aggregate_generation_availability(
    mappings: Iterable[dict[str, str]], values: dict[str, object]
) -> dict[str, float]:
    weighted: dict[str, float] = {}
    capacities: dict[str, float] = {}
    for mapping in mappings:
        if mapping["included"].lower() != "true":
            continue
        capacity_query = mapping["query_key"]
        capacity = _future_value(values, capacity_query)
        if capacity <= 0.0:
            continue
        availability_query = capacity_query.replace(
            "_capacity_in_merit_order_table",
            "_availability_in_merit_order_table",
        )
        group = mapping["asset_group"]
        weighted[group] = weighted.get(group, 0.0) + capacity * _future_value(
            values, availability_query
        ) / 100.0
        capacities[group] = capacities.get(group, 0.0) + capacity
    return {group: value / capacities[group] for group, value in weighted.items()}


def load_generation_parameters(
    scenario_key: str,
    year: int,
    scenario_inventory_path: Path,
    aggregation_path: Path,
) -> tuple[dict[str, float], dict[str, float]]:
    inventory_rows = load_case_rows(scenario_inventory_path, scenario_key, year)
    if len(inventory_rows) != 1:
        raise SourceValidationError(
            f"Expected one scenario inventory row for {scenario_key} {year}."
        )
    with aggregation_path.open(encoding="utf-8-sig", newline="") as input_file:
        mappings = list(csv.DictReader(input_file))
    relevant = [
        row
        for row in mappings
        if row["included"].lower() == "true"
    ]
    queries = [
        query
        for row in relevant
        for query in (
            row["query_key"],
            row["query_key"].replace(
                "_capacity_in_merit_order_table",
                "_operating_costs_in_merit_order_table",
            ),
            row["query_key"].replace(
                "_capacity_in_merit_order_table",
                "_availability_in_merit_order_table",
            ),
        )
    ]
    inventory = inventory_rows[0]
    with EtmClient(str(inventory["engine_base_url"])) as client:
        response = client.query_scenario(int(inventory["scenario_id"]), queries)
    return (
        aggregate_generation_costs(relevant, response.values),
        aggregate_generation_availability(relevant, response.values),
    )


def load_technology_parameters(
    scenario_key: str,
    year: int,
    hydrogen_rows: Iterable[dict[str, object]],
    generation_costs: dict[str, float],
    scenario_inventory_path: Path,
    electrolysis_efficiency: float | None,
    i_elgas_technology_data_path: Path,
) -> tuple[
    dict[str, TechnologyParameter], dict[str, float], dict[str, float]
]:
    inventory_rows = load_case_rows(scenario_inventory_path, scenario_key, year)
    if len(inventory_rows) != 1:
        raise SourceValidationError(
            f"Expected one scenario inventory row for {scenario_key} {year}."
        )
    inventory = inventory_rows[0]
    active_hydrogen_groups = {
        str(row["asset_group"])
        for row in hydrogen_rows
        if float(row["capacity_mw"]) > 0.0
    }
    energy_queries = [
        query
        for group, pair in HYDROGEN_CONVERSION_ENERGY_QUERIES.items()
        if group in active_hydrogen_groups
        for query in pair
    ]
    parameter_queries = energy_queries + list(NATURAL_GAS_CHP_EFFICIENCY_INPUTS)
    with EtmClient(str(inventory["engine_base_url"])) as client:
        scenario = client.get_scenario(int(inventory["scenario_id"]))
        response = (
            client.query_scenario(int(inventory["scenario_id"]), parameter_queries)
            if parameter_queries
            else None
        )
    user_values = scenario.get("user_values", {})
    if not isinstance(user_values, dict):
        raise SourceValidationError("ETM scenario user values are not a mapping.")

    fuel_prices = {
        "methane": float(user_values["costs_gas"]),
        "hydrogen": float(user_values["costs_hydrogen"]),
        "renewable_methane": float(user_values["costs_greengas"]),
    }
    co2_price = float(user_values["costs_co2"])
    i_elgas_parameters = load_i_elgas_technology_parameters(
        i_elgas_technology_data_path
    )
    import_efficiencies = {
        **HYDROGEN_CARRIER_IMPORT_EFFICIENCIES,
        "ammonia_reformer": float(user_values["efficiency_ammonia_reforming"])
        / 100.0,
    }
    hydrogen_carrier_import_costs = {
        group: float(user_values[input_key]) / import_efficiencies[group]
        for group, input_key in HYDROGEN_CARRIER_IMPORT_COST_INPUTS.items()
    }
    parameters: dict[str, TechnologyParameter] = {}
    for group, input_key in POWER_CONVERSION_EFFICIENCY_INPUTS.items():
        if group not in generation_costs:
            continue
        efficiency = float(user_values[input_key]) / 100.0
        fuel = "hydrogen" if group.startswith("hydrogen_") else "methane"
        residual_cost = generation_costs[group] - fuel_prices[fuel] / efficiency
        parameters[group] = TechnologyParameter(
            group,
            efficiency,
            max(0.0, residual_cost),
            f"ETM scenario user_values.{input_key}",
            "ETM merit-order marginal cost minus upstream fuel price / efficiency",
        )

    chp_capacities = {
        query: _future_value(response.values, query)
        for query in NATURAL_GAS_CHP_EFFICIENCY_INPUTS
    }
    total_chp_capacity = sum(chp_capacities.values())
    if total_chp_capacity > 0.0 and "natural_gas_chp" in generation_costs:
        efficiency = sum(
            capacity
            * float(user_values[NATURAL_GAS_CHP_EFFICIENCY_INPUTS[query]])
            / 100.0
            for query, capacity in chp_capacities.items()
        ) / total_chp_capacity
        parameters["natural_gas_chp"] = TechnologyParameter(
            "natural_gas_chp",
            efficiency,
            max(
                0.0,
                generation_costs["natural_gas_chp"]
                - fuel_prices["methane"] / efficiency,
            ),
            "capacity-weighted ETM electrical efficiencies; useful heat omitted",
            "ETM merit-order marginal cost minus upstream fuel price / efficiency",
        )

    if electrolysis_efficiency is not None:
        parameters["electrolysis"] = TechnologyParameter(
            "electrolysis",
            electrolysis_efficiency,
            0.0,
            "ETM hydrogen-output MW / regional workbook electricity-input MW",
            "zero variable O&M; electricity input is priced endogenously",
        )
    values = response.values if response is not None else {}
    for group, (input_query, output_query) in HYDROGEN_CONVERSION_ENERGY_QUERIES.items():
        if group not in active_hydrogen_groups:
            continue
        input_energy = _future_value(values, input_query)
        output_energy = _future_value(values, output_query)
        if input_energy <= 0.0 or output_energy <= 0.0:
            raise SourceValidationError(
                f"Cannot derive {group} efficiency for {scenario_key} {year}: "
                f"input={input_energy}, output={output_energy}."
            )
        parameters[group] = TechnologyParameter(
            group,
            output_energy / input_energy,
            0.0,
            f"ETM {output_query} / {input_query}",
            "zero unresolved variable O&M; methane input is priced upstream",
        )
    for group, source_technology in I_ELGAS_REFORMER_MAPPING.items():
        if group not in active_hydrogen_groups:
            continue
        source = i_elgas_parameters[source_technology]
        current = parameters[group]
        co2_cost = reformer_co2_cost_eur_per_mwh_output(
            source.emissions_kg_per_gj_input,
            current.conversion_efficiency,
            co2_price,
        )
        parameters[group] = TechnologyParameter(
            group,
            current.conversion_efficiency,
            source.marginal_cost_eur_per_mwh_output + co2_cost,
            current.efficiency_source,
            "I-ELGAS non-fuel marginal cost plus residual emissions at ETM CO2 price; methane priced upstream",
            source.marginal_cost_eur_per_mwh_output,
            source.emissions_kg_per_gj_input,
            co2_cost,
            co2_price,
        )
    return parameters, fuel_prices, hydrogen_carrier_import_costs


def load_i_elgas_technology_parameters(
    path: Path,
) -> dict[str, IElgasTechnologyParameter]:
    parameters: dict[str, IElgasTechnologyParameter] = {}
    with path.open(encoding="utf-8-sig", newline="") as input_file:
        for row in csv.DictReader(input_file, delimiter=";"):
            technology = row["TECH/PROCESS"].strip()
            if technology not in set(I_ELGAS_REFORMER_MAPPING.values()):
                continue
            if technology in parameters:
                raise SourceValidationError(
                    f"Duplicate I-ELGAS technology row: {technology}."
                )
            parameters[technology] = IElgasTechnologyParameter(
                technology,
                float(row["EMISSIONS"].replace(",", ".")),
                float(row["Marginal Costs"].replace(",", ".")),
            )
    missing = set(I_ELGAS_REFORMER_MAPPING.values()) - set(parameters)
    if missing:
        raise SourceValidationError(
            f"Missing I-ELGAS reformer technologies: {sorted(missing)}."
        )
    return parameters


def build_static_availability_profile_tables(
    electricity_rows: Iterable[dict[str, object]],
    year: int,
    availability: dict[str, float],
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    excluded = set(TYNDP_PROFILE_SOURCES) | set(MUST_RUN_CURVE_SOURCES)
    rows = [
        row
        for row in electricity_rows
        if float(row["capacity_mw"]) > 0.0
        and str(row["asset_group"]) in availability
        and str(row["asset_group"]) not in excluded
    ]
    groups = sorted({str(row["asset_group"]) for row in rows})
    attachments = [
        {
            "asset": f'{row["node"]}_{row["asset_group"]}_{row["operating_mode"]}',
            "commission_year": year,
            "profile_name": f'SR2025_{row["asset_group"]}_{year}',
            "profile_type": "availability",
        }
        for row in rows
    ]
    profiles = [
        {
            "milestone_year": year,
            "profile_name": f"SR2025_{group}_{year}",
            "rep_period": 1,
            "timestep": timestep,
            "value": round(availability[group], 6),
        }
        for group in groups
        for timestep in range(1, 8761)
    ]
    return attachments, profiles


def build_renewable_profile_tables(
    electricity_rows: Iterable[dict[str, object]],
    year: int,
    tyndp_dir: Path,
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    electricity_rows = list(electricity_rows)
    source_assets: dict[str, float] = {}
    with (tyndp_dir / "asset.csv").open(encoding="utf-8-sig", newline="") as input_file:
        for row in csv.DictReader(input_file):
            if row["asset"] in {
                asset
                for assets in TYNDP_PROFILE_SOURCES.values()
                for asset in assets
            }:
                source_assets[row["asset"]] = float(row["capacity"])

    required_sources = {
        asset
        for group, assets in TYNDP_PROFILE_SOURCES.items()
        if any(
            str(row["asset_group"]) == group and float(row["capacity_mw"]) > 0.0
            for row in electricity_rows
        )
        for asset in assets
    }
    missing = required_sources - source_assets.keys()
    if missing:
        raise SourceValidationError(
            f"Missing TYNDP renewable profile source assets: {sorted(missing)}"
        )

    source_values: dict[tuple[str, str, str], dict[str, float]] = {}
    with (tyndp_dir / "profiles-rep-periods.csv").open(
        encoding="utf-8-sig", newline=""
    ) as input_file:
        for row in csv.DictReader(input_file):
            profile_name = row["profile_name"]
            if profile_name not in required_sources:
                continue
            key = (row["rep_period"], row["timestep"], row["milestone_year"])
            source_values.setdefault(key, {})[profile_name] = float(row["value"])

    attachments: list[dict[str, object]] = []
    active_groups: set[str] = set()
    for row in electricity_rows:
        group = str(row["asset_group"])
        if group not in TYNDP_PROFILE_SOURCES or float(row["capacity_mw"]) <= 0.0:
            continue
        active_groups.add(group)
        attachments.append({
            "asset": f'{row["node"]}_{group}_{row["operating_mode"]}',
            "commission_year": year,
            "profile_name": f"SR2025_{group}_{year}",
            "profile_type": "availability",
        })

    profiles: list[dict[str, object]] = []
    for rep_period, timestep, _ in sorted(
        source_values, key=lambda key: (int(key[0]), int(key[1]))
    ):
        values = source_values[(rep_period, timestep, _)]
        for group in sorted(active_groups):
            sources = TYNDP_PROFILE_SOURCES[group]
            denominator = sum(source_assets[source] for source in sources)
            if denominator <= 0.0 or not all(source in values for source in sources):
                raise SourceValidationError(
                    f"Incomplete TYNDP profile values for {group}, rep period "
                    f"{rep_period}, timestep {timestep}."
                )
            value = sum(
                source_assets[source] * values[source] for source in sources
            ) / denominator
            profiles.append({
                "milestone_year": year,
                "profile_name": f"SR2025_{group}_{year}",
                "rep_period": rep_period,
                "timestep": timestep,
                "value": round(value, 6),
            })
    return attachments, profiles


def build_must_run_profile_tables(
    electricity_rows: Iterable[dict[str, object]],
    year: int,
    curve_rows: Iterable[dict[str, object]],
) -> tuple[list[dict[str, object]], list[dict[str, object]], dict[str, float]]:
    electricity_rows = list(electricity_rows)
    group_capacities = {
        group: sum(
            float(row["capacity_mw"])
            for row in electricity_rows
            if str(row["asset_group"]) == group
        )
        for group in MUST_RUN_CURVE_SOURCES
    }
    active_groups = {
        group for group, capacity in group_capacities.items() if capacity > 0.0
    }
    attachments = [
        {
            "asset": f'{row["node"]}_{row["asset_group"]}_{row["operating_mode"]}',
            "commission_year": year,
            "profile_name": f'SR2025_{row["asset_group"]}_{year}',
            "profile_type": "availability",
        }
        for row in electricity_rows
        if str(row["asset_group"]) in active_groups
        and float(row["capacity_mw"]) > 0.0
    ]
    profiles: list[dict[str, object]] = []
    profile_sums = {group: 0.0 for group in active_groups}
    for timestep, curve_row in enumerate(curve_rows, start=1):
        for group in sorted(active_groups):
            output = sum(
                float(curve_row[f"{source}.output (MW)"] or 0.0)
                for source in MUST_RUN_CURVE_SOURCES[group]
            )
            value = output / group_capacities[group]
            if value < -1e-9 or value > 1.0 + 1e-6:
                raise SourceValidationError(
                    f"ETM must-run profile for {group} is outside [0, 1] at "
                    f"timestep {timestep}: {value}."
                )
            value = round(min(1.0, max(0.0, value)), 6)
            profile_sums[group] += value
            profiles.append({
                "milestone_year": year,
                "profile_name": f"SR2025_{group}_{year}",
                "rep_period": 1,
                "timestep": timestep,
                "value": value,
            })
    minimum_energy = {
        f'{row["node"]}_{group}_{row["operating_mode"]}': (
            float(row["capacity_mw"]) * profile_sums[group]
        )
        for row in electricity_rows
        if (group := str(row["asset_group"])) in active_groups
    }
    return attachments, profiles, minimum_energy


def load_must_run_profile_tables(
    scenario_key: str,
    year: int,
    electricity_rows: Iterable[dict[str, object]],
    scenario_inventory_path: Path,
) -> tuple[list[dict[str, object]], list[dict[str, object]], dict[str, float]]:
    inventory_rows = load_case_rows(scenario_inventory_path, scenario_key, year)
    if len(inventory_rows) != 1:
        raise SourceValidationError(
            f"Expected one scenario inventory row for {scenario_key} {year}."
        )
    inventory = inventory_rows[0]
    with EtmClient(str(inventory["engine_base_url"])) as client:
        curve = client.get_curve_csv(int(inventory["scenario_id"]), "merit_order")
    return build_must_run_profile_tables(electricity_rows, year, curve.rows)


def load_hydrogen_import_costs(
    scenario_key: str, year: int, hydrogen_capacity_scan_path: Path
) -> dict[str, float]:
    query_groups = {
        "energy_imported_hydrogen_backup_h2_chart": "hydrogen_import_backup",
        "energy_imported_hydrogen_baseload_h2_chart": "hydrogen_import_baseload",
    }
    costs: dict[str, float] = {}
    with hydrogen_capacity_scan_path.open(
        encoding="utf-8-sig", newline=""
    ) as input_file:
        for row in csv.DictReader(input_file):
            group = query_groups.get(row["query_key"])
            if (
                group is not None
                and row["scenario_key"] == scenario_key
                and int(row["year"]) == year
                and float(row["capacity_mw_hydrogen"]) > 0.0
            ):
                costs[group] = float(row["operating_cost_eur_per_mwh"])
    return costs


def load_hydrogen_carrier_production(
    scenario_key: str, year: int, hydrogen_capacity_scan_path: Path
) -> dict[tuple[str, str], float]:
    query_modes = {
        "energy_hydrogen_ammonia_reformer_dispatchable_h2_chart": (
            "ammonia_reformer", "dispatchable"
        ),
        "energy_hydrogen_ammonia_reformer_must_run_h2_chart": (
            "ammonia_reformer", "must_run"
        ),
        "energy_hydrogen_liquid_hydrogen_regasifier_h2_chart": (
            "liquid_hydrogen_regasifier", "unspecified"
        ),
        "energy_hydrogen_lohc_reformer_h2_chart": (
            "lohc_reformer", "unspecified"
        ),
    }
    production: dict[tuple[str, str], float] = {}
    with hydrogen_capacity_scan_path.open(
        encoding="utf-8-sig", newline=""
    ) as input_file:
        for row in csv.DictReader(input_file):
            key = query_modes.get(row["query_key"])
            if (
                key is not None
                and row["scenario_key"] == scenario_key
                and int(row["year"]) == year
                and float(row["capacity_mw_hydrogen"]) > 0.0
            ):
                production[key] = float(row["production_twh"])
    return production


def load_storage_technologies(
    scenario_key: str,
    year: int,
    hydrogen_rows: Iterable[dict[str, object]],
    scenario_inventory_path: Path,
) -> list[StorageTechnologyRow]:
    inventory_rows = load_case_rows(scenario_inventory_path, scenario_key, year)
    if len(inventory_rows) != 1:
        raise SourceValidationError(
            f"Expected one scenario inventory row for {scenario_key} {year}."
        )
    inventory = inventory_rows[0]
    standard_queries = [
        f"{prefix}_{suffix}"
        for prefix in ELECTRICITY_STORAGE_QUERIES.values()
        for suffix in ("input_capacity", "output_capacity", "storage_volume", "efficiency")
    ]
    coupled_capacity_queries = [
        "merit_order_battery_solar_pv_capacity_in_merit_order_table",
        "merit_order_battery_wind_inland_capacity_in_merit_order_table",
    ]
    hydrogen_volume_queries = [
        "hydrogen_storage_volume_salt_cavern_in_mekko_of_storage_volume",
        "hydrogen_storage_volume_depleted_gas_field_in_mekko_of_storage_volume",
    ]
    with EtmClient(str(inventory["engine_base_url"])) as client:
        response = client.query_scenario(
            int(inventory["scenario_id"]),
            standard_queries + coupled_capacity_queries + hydrogen_volume_queries,
        )
        scenario = client.get_scenario(int(inventory["scenario_id"]))
    values = response.values
    user_values = scenario.get("user_values", {})
    if not isinstance(user_values, dict):
        raise SourceValidationError("ETM scenario user values are not a mapping.")

    rows: list[StorageTechnologyRow] = []
    for group, prefix in ELECTRICITY_STORAGE_QUERIES.items():
        input_capacity = _future_value(values, f"{prefix}_input_capacity")
        output_capacity = _future_value(values, f"{prefix}_output_capacity")
        volume = _future_value(values, f"{prefix}_storage_volume")
        efficiency = _future_value(values, f"{prefix}_efficiency") / 100.0
        if group == "battery_system":
            solar_power = _future_value(
                values, "merit_order_battery_solar_pv_capacity_in_merit_order_table"
            ) * float(user_values["battery_capacity_always_on_solar_pv_solar_radiation"]) / 100.0
            wind_power = _future_value(
                values, "merit_order_battery_wind_inland_capacity_in_merit_order_table"
            ) * float(user_values["battery_capacity_always_on_wind_turbine_inland"]) / 100.0
            input_capacity += solar_power + wind_power
            output_capacity += solar_power + wind_power
            volume += (
                solar_power
                * float(user_values["volume_of_energy_flexibility_solar_batteries_electricity"])
                + wind_power
                * float(user_values["volume_of_energy_flexibility_wind_batteries_electricity"])
            )
        rows.append(StorageTechnologyRow(
            group,
            "electricity",
            input_capacity,
            output_capacity,
            volume,
            1.0,
            efficiency,
            False,
        ))

    hydrogen_capacities = {
        str(row["asset_group"]): float(row["capacity_mw"])
        for row in hydrogen_rows
        if str(row["asset_group"]).startswith("hydrogen_storage_")
    }
    for kind in ("salt_cavern", "depleted_gas_field"):
        group = f"hydrogen_storage_{kind}"
        capacity = hydrogen_capacities.get(group, 0.0)
        volume = _future_value(
            values, f"hydrogen_storage_volume_{kind}_in_mekko_of_storage_volume"
        ) * 1_000_000.0
        rows.append(StorageTechnologyRow(
            group, "hydrogen", capacity, capacity, volume, 1.0, 1.0, True
        ))
    return rows


def calculate_electrolysis_efficiency(
    scenario_key: str,
    year: int,
    hydrogen_rows: Iterable[dict[str, object]],
    workbook_path: Path,
) -> ConversionEfficiencyRow | None:
    """Derive the output/input capacity ratio used by the Tulipa conversion."""
    output_capacity_mw = sum(
        float(row["capacity_mw"])
        for row in hydrogen_rows
        if row["asset_group"] == "electrolysis"
    )
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    input_capacity_mw = 0.0
    for sheet_name in (
        "Municipality (ELEC, capacity)",
        "Industry (ELEC, capacity)",
    ):
        rows = workbook[sheet_name].iter_rows(values_only=True)
        metadata = [next(rows) for _ in range(6)]
        sectors = metadata[4]
        indices = [
            index
            for index, sector in enumerate(sectors)
            if sector in {"Power_to_gas_onshore", "Power_to_gas_offshore"}
        ]
        input_capacity_mw += sum(
            abs(float(values[index] or 0.0))
            for values in rows
            for index in indices
        )
    if output_capacity_mw == 0.0:
        return None
    if input_capacity_mw <= 0.0:
        raise SourceValidationError(
            f"Incomplete electrolysis capacity basis for {scenario_key} {year}: "
            f"input={input_capacity_mw}, output={output_capacity_mw}."
        )
    efficiency = output_capacity_mw / input_capacity_mw
    if not 0.0 < efficiency <= 1.0:
        raise SourceValidationError(
            f"Invalid electrolysis efficiency for {scenario_key} {year}: {efficiency}."
        )
    return ConversionEfficiencyRow(
        scenario_key,
        year,
        "electrolysis",
        input_capacity_mw,
        output_capacity_mw,
        efficiency,
        "asset capacity is MW_H2 output; efficiency converts MW_e input",
    )


def write_table(path: Path, columns: tuple[str, ...], rows: Iterable[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as output_file:
        writer = csv.DictWriter(output_file, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def load_demand_profile_tables(
    scenario_key: str,
    year: int,
    electricity_nodes: Iterable[str],
    electricity_demand_path: Path,
    grouped_demand_path: Path,
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    """Build Tulipa demand attachments from validated hourly demand outputs."""
    values: dict[str, dict[int, float]] = {
        electricity_bus(node, year): {} for node in electricity_nodes
    }
    with electricity_demand_path.open(encoding="utf-8-sig", newline="") as input_file:
        for row in csv.DictReader(input_file):
            if row["scenario_key"] != scenario_key or int(row["year"]) != year:
                continue
            asset = electricity_bus(str(row["node"]), year)
            if asset in values:
                values[asset][int(row["hour"])] = float(row["demand_mw"])

    carrier_assets = {
        "hydrogen": country_bus("hydrogen", year),
        "methane": country_bus("methane", year),
    }
    values.update({asset: {} for asset in carrier_assets.values()})
    with grouped_demand_path.open(encoding="utf-8-sig", newline="") as input_file:
        for row in csv.DictReader(input_file):
            if row["scenario_key"] != scenario_key or int(row["year"]) != year:
                continue
            asset = carrier_assets.get(row["carrier"])
            if asset is None:
                continue
            hour = int(row["hour"])
            values[asset][hour] = values[asset].get(hour, 0.0) + float(
                row["normalized_mw"]
            )

    expected_hours = set(range(1, HOURS_PER_YEAR + 1))
    incomplete = {
        asset: len(hourly)
        for asset, hourly in values.items()
        if set(hourly) != expected_hours
    }
    if incomplete:
        raise SourceValidationError(
            f"Incomplete demand profiles for {scenario_key} {year}: {incomplete}."
        )

    attachments: list[dict[str, object]] = []
    profiles: list[dict[str, object]] = []
    for asset, hourly in sorted(values.items()):
        profile_name = f"SR2025_{asset}"
        attachments.append({
            "asset": asset,
            "commission_year": year,
            "profile_name": profile_name,
            "profile_type": "demand",
        })
        profiles.extend({
            "milestone_year": year,
            "profile_name": profile_name,
            "rep_period": 1,
            "timestep": hour,
            "value": hourly[hour],
        } for hour in range(1, HOURS_PER_YEAR + 1))
    return attachments, profiles


def export_case(
    scenario_key: str,
    year: int,
    electricity_input: Path,
    hydrogen_input: Path,
    regional_dir: Path,
    output_root: Path,
    trading_capacities_input: Path = Path(
        "source_data/i_elgas/Electricity Trading Capacities I-ELGAS.xlsx"
    ),
    scenario_inventory_input: Path = Path("output/audit/scenario_inventory.csv"),
    electricity_aggregation_input: Path = Path(
        "config/electricity_capacity_aggregation.csv"
    ),
    tyndp_root: Path = Path("../TYNDP-26-to-Tulipa"),
    hydrogen_capacity_scan_input: Path = Path(
        "output/audit/hydrogen_capacity_all_scenarios.csv"
    ),
    methane_supply_input: Path = Path("output/audit/methane_supply_grouped.csv"),
    electricity_demand_input: Path = Path(
        "output/regionalisation/electricity_demand_node_hourly.csv"
    ),
    grouped_demand_input: Path = Path("output/profiles/demand_hourly.csv"),
    i_elgas_technology_data_input: Path = Path(
        "source_data/i_elgas/I-ELGAS_Technology_Data.csv"
    ),
    renewable_profile_input: Path | None = None,
) -> Path:
    electricity_rows = load_case_rows(electricity_input, scenario_key, year)
    hydrogen_rows = load_case_rows(hydrogen_input, scenario_key, year)
    methane_rows = load_case_rows(methane_supply_input, scenario_key, year)
    reference = load_regional_reference(scenario_key, year, regional_dir)
    scenario_name = SCENARIO_NAMES[scenario_key]
    reference_year = max(2030, min(2050, year))
    workbook_path = (
        regional_dir
        / f"Scenario {scenario_name}"
        / f"Scenario {scenario_name} {reference_year}.xlsx"
    )
    efficiency_row = calculate_electrolysis_efficiency(
        scenario_key, year, hydrogen_rows, workbook_path
    )
    weights, electricity_nodes = load_hydrogen_node_weights(
        scenario_key, year, regional_dir
    )
    storage_rows = load_storage_technologies(
        scenario_key, year, hydrogen_rows, scenario_inventory_input
    )
    electricity_storage_weights = load_electricity_storage_weights(
        scenario_key, year, regional_dir
    )
    generation_costs, generation_availability = load_generation_parameters(
        scenario_key,
        year,
        scenario_inventory_input,
        electricity_aggregation_input,
    )
    technology_parameters, fuel_prices, hydrogen_carrier_import_costs = (
        load_technology_parameters(
        scenario_key,
        year,
        hydrogen_rows,
        generation_costs,
        scenario_inventory_input,
        efficiency_row.conversion_efficiency if efficiency_row is not None else None,
        i_elgas_technology_data_input,
        )
    )
    hydrogen_import_costs = load_hydrogen_import_costs(
        scenario_key, year, hydrogen_capacity_scan_input
    )
    hydrogen_carrier_production = load_hydrogen_carrier_production(
        scenario_key, year, hydrogen_capacity_scan_input
    )
    transport_links = load_electricity_transport_links(
        trading_capacities_input, year
    )
    frames, reconciliation, storage_reconciliation = assemble_tulipa_tables(
        year,
        electricity_rows,
        hydrogen_rows,
        weights,
        electricity_nodes,
        transport_links,
        technology_parameters,
        storage_rows,
        electricity_storage_weights,
        generation_costs,
        hydrogen_carrier_import_costs,
        hydrogen_import_costs,
        hydrogen_carrier_production,
        methane_rows,
        fuel_prices,
    )
    output_dir = output_root / f"{scenario_key}_{year}"
    if renewable_profile_input is None:
        profile_dir_name = TYNDP_PROFILE_DIRECTORIES.get(year)
        if profile_dir_name is None:
            raise SourceValidationError(
                f"No TYNDP renewable profile year mapping for {year}."
            )
        renewable_profile_input = tyndp_root / profile_dir_name
    profile_attachments, renewable_profiles = build_renewable_profile_tables(
        electricity_rows, year, renewable_profile_input
    )
    must_run_attachments, must_run_profiles, minimum_energy = (
        load_must_run_profile_tables(
            scenario_key, year, electricity_rows, scenario_inventory_input
        )
    )
    static_attachments, static_profiles = build_static_availability_profile_tables(
        electricity_rows, year, generation_availability
    )
    demand_attachments, demand_profiles = load_demand_profile_tables(
        scenario_key,
        year,
        electricity_nodes,
        electricity_demand_input,
        grouped_demand_input,
    )
    for row in frames["asset-milestone"]:
        if row["asset"] in minimum_energy:
            row["min_energy_timeframe_partition"] = minimum_energy[row["asset"]]
    for table, columns in TABLE_COLUMNS.items():
        write_table(output_dir / f"{table}.csv", columns, frames[table])
    write_table(
        output_dir / "assets-profiles.csv",
        ("asset", "commission_year", "profile_name", "profile_type"),
        profile_attachments
        + must_run_attachments
        + static_attachments
        + demand_attachments,
    )
    write_table(
        output_dir / "profiles-rep-periods.csv",
        ("milestone_year", "profile_name", "rep_period", "timestep", "value"),
        renewable_profiles + must_run_profiles + static_profiles + demand_profiles,
    )
    reconciliation_columns = (
        "asset_group", "operating_mode", "national_capacity_mw",
        "distributed_capacity_mw", "difference_mw",
    )
    write_table(
        output_dir / "conversion-capacity-reconciliation.csv",
        reconciliation_columns,
        reconciliation,
    )
    emitted_assets = {row["asset"]: float(row["capacity"]) for row in frames["asset"]}
    generation_groups = sorted({
        (str(row["asset_group"]), str(row["operating_mode"]))
        for row in electricity_rows
    })
    generation_reconciliation = []
    for group, mode in generation_groups:
        source_rows = [
            row
            for row in electricity_rows
            if str(row["asset_group"]) == group
            and str(row["operating_mode"]) == mode
        ]
        source_capacity = sum(float(row["capacity_mw"]) for row in source_rows)
        tulipa_capacity = sum(
            emitted_assets.get(f'{row["node"]}_{group}_{mode}', 0.0)
            for row in source_rows
        )
        generation_reconciliation.append({
            "asset_group": group,
            "operating_mode": mode,
            "source_capacity_mw": source_capacity,
            "tulipa_capacity_mw": tulipa_capacity,
            "difference_mw": tulipa_capacity - source_capacity,
        })
    write_table(
        output_dir / "generation-capacity-reconciliation.csv",
        (
            "asset_group", "operating_mode", "source_capacity_mw",
            "tulipa_capacity_mw", "difference_mw",
        ),
        generation_reconciliation,
    )
    write_table(
        output_dir / "conversion-efficiency-audit.csv",
        (
            "scenario_key", "year", "asset_group", "input_capacity_mw",
            "output_capacity_mw", "conversion_efficiency", "capacity_basis",
        ),
        [efficiency_row.__dict__] if efficiency_row is not None else [],
    )
    write_table(
        output_dir / "technology-parameters.csv",
        (
            "asset_group", "conversion_efficiency",
            "variable_cost_eur_per_mwh_output", "efficiency_source", "cost_source",
            "non_fuel_variable_cost_eur_per_mwh_output",
            "emissions_kg_per_gj_input", "co2_cost_eur_per_mwh_output",
            "co2_price_eur_per_tonne",
        ),
        [row.__dict__ for row in sorted(
            technology_parameters.values(), key=lambda item: item.asset_group
        )],
    )
    write_table(
        output_dir / "hydrogen-import-parameters.csv",
        ("asset_group", "variable_cost_eur_per_mwh_hydrogen", "cost_source"),
        [
            {
                "asset_group": group,
                "variable_cost_eur_per_mwh_hydrogen": cost,
                "cost_source": "ETM imported-carrier price / conversion yield",
            }
            for group, cost in sorted(hydrogen_carrier_import_costs.items())
        ],
    )
    write_table(
        output_dir / "storage-capacity-reconciliation.csv",
        (
            "asset_group", "carrier", "source_input_capacity_mw",
            "tulipa_input_capacity_mw", "input_difference_mw",
            "source_output_capacity_mw", "tulipa_output_capacity_mw",
            "output_difference_mw", "source_storage_volume_mwh",
            "tulipa_storage_volume_mwh", "storage_volume_difference_mwh",
        ),
        storage_reconciliation,
    )
    write_table(
        output_dir / "electricity-transport-capacity-audit.csv",
        (
            "from_endpoint", "to_endpoint", "forward_capacity_mw",
            "reverse_capacity_mw", "source_year", "link_scope",
        ),
        transport_links,
    )
    (output_dir / "provenance.txt").write_text(
        f"scenario: {scenario_key}\n"
        f"year: {year}\n"
        f"electricity_capacity_source: {electricity_input}\n"
        f"hydrogen_capacity_source: {hydrogen_input}\n"
        f"methane_supply_source: {methane_supply_input}\n"
        f"electricity_demand_source: {electricity_demand_input}\n"
        f"grouped_demand_source: {grouped_demand_input}\n"
        f"i_elgas_technology_data_source: {i_elgas_technology_data_input}\n"
        f"electricity_transport_source: {trading_capacities_input}\n"
        f"renewable_profile_source: {renewable_profile_input}\n"
        f"storage_parameter_source: {scenario_inventory_input}\n"
        f"hydrogen_import_parameter_source: {hydrogen_capacity_scan_input}\n"
        f"methane_price_eur_per_mwh: {fuel_prices['methane']}\n"
        f"hydrogen_price_eur_per_mwh: {fuel_prices['hydrogen']}\n"
        f"tulipa_schema: v{TULIPA_SCHEMA_VERSION} ({TULIPA_SCHEMA_URL})\n"
        f"electricity_transport_source_year: {2027 if year == 2025 else year}\n"
        "electricity_nodes: authoritative I-ELGAS E-node topology\n"
        "hydrogen_node: NL_H_Demand_<year>\n"
        "methane_node: NL_M_Demand_<year>\n"
        "converter_placement: municipality H2-capacity drivers aggregated to E-nodes\n"
        "scope: generation, conversions, electricity storage, hydrogen storage, and electricity transport\n"
        "renewable_profiles: Dutch TYNDP PECD v4.2 profiles shared by SR technology group\n"
        "must_run_profiles: ETM merit-order output with matching annual minimum energy\n"
        "generation_costs: capacity-weighted ETM merit-order marginal costs\n"
        "electrolysis_capacity_basis: MW_H2 output\n"
        "electrolysis_balance: conversion_efficiency * MW_e input = MW_H2 output\n"
        "thermal_generation_boundary: direct MW-electric producers; CHP is electricity-only\n",
        encoding="utf-8",
    )
    return output_dir


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Generate Tulipa node and coupled conversion-asset tables from cached outputs."
    )
    parser.add_argument("scenario_key", choices=tuple(SCENARIO_NAMES))
    parser.add_argument("year", type=int)
    parser.add_argument(
        "--electricity-input",
        type=Path,
        default=Path("output/regionalisation/electricity_capacity_by_node.csv"),
    )
    parser.add_argument(
        "--hydrogen-input",
        type=Path,
        default=Path("output/audit/hydrogen_capacity_grouped.csv"),
    )
    parser.add_argument(
        "--methane-supply-input",
        type=Path,
        default=Path("output/audit/methane_supply_grouped.csv"),
    )
    parser.add_argument(
        "--electricity-demand-input",
        type=Path,
        default=Path("output/regionalisation/electricity_demand_node_hourly.csv"),
    )
    parser.add_argument(
        "--grouped-demand-input",
        type=Path,
        default=Path("output/profiles/demand_hourly.csv"),
    )
    parser.add_argument(
        "--i-elgas-technology-data-input",
        type=Path,
        default=Path("source_data/i_elgas/I-ELGAS_Technology_Data.csv"),
    )
    parser.add_argument(
        "--regional-dir", type=Path, default=Path("source_data/regionalisation")
    )
    parser.add_argument(
        "--trading-capacities-input",
        type=Path,
        default=Path(
            "source_data/i_elgas/Electricity Trading Capacities I-ELGAS.xlsx"
        ),
    )
    parser.add_argument("--output-root", type=Path, default=Path("output/tulipa"))
    parser.add_argument(
        "--scenario-inventory-input",
        type=Path,
        default=Path("output/audit/scenario_inventory.csv"),
    )
    parser.add_argument(
        "--electricity-aggregation-input",
        type=Path,
        default=Path("config/electricity_capacity_aggregation.csv"),
    )
    parser.add_argument(
        "--tyndp-root",
        type=Path,
        default=Path("../TYNDP-26-to-Tulipa"),
    )
    parser.add_argument(
        "--renewable-profile-input",
        type=Path,
        help="Exact TYNDP dispatch folder supplying Dutch VRE profiles.",
    )
    args = parser.parse_args()
    output_dir = export_case(
        args.scenario_key,
        args.year,
        args.electricity_input,
        args.hydrogen_input,
        args.regional_dir,
        args.output_root,
        args.trading_capacities_input,
        args.scenario_inventory_input,
        electricity_aggregation_input=args.electricity_aggregation_input,
        tyndp_root=args.tyndp_root,
        renewable_profile_input=args.renewable_profile_input,
        methane_supply_input=args.methane_supply_input,
        electricity_demand_input=args.electricity_demand_input,
        grouped_demand_input=args.grouped_demand_input,
        i_elgas_technology_data_input=args.i_elgas_technology_data_input,
    )
    print(f"Created Tulipa asset and flow tables in {output_dir}.")


if __name__ == "__main__":
    main()