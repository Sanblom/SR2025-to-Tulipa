import pytest
from openpyxl import Workbook

from sr2025_to_tulipa.tulipa_assets import (
    StorageTechnologyRow,
    TABLE_COLUMNS,
    TechnologyParameter,
    aggregate_generation_availability,
    aggregate_generation_costs,
    assemble_tulipa_tables,
    build_renewable_profile_tables,
    build_must_run_profile_tables,
    build_static_availability_profile_tables,
    calculate_electrolysis_efficiency,
    country_bus,
    electricity_bus,
    load_demand_profile_tables,
    load_electricity_transport_links,
    load_i_elgas_technology_parameters,
    reformer_co2_cost_eur_per_mwh_output,
)


def test_generation_costs_are_capacity_weighted_at_group_boundary() -> None:
    mappings = [
        {
            "query_key": "source_a_capacity_in_merit_order_table",
            "asset_group": "solar_pv",
            "fuel": "solar",
            "included": "true",
        },
        {
            "query_key": "source_b_capacity_in_merit_order_table",
            "asset_group": "solar_pv",
            "fuel": "solar",
            "included": "true",
        },
        {
            "query_key": "source_c_capacity_in_merit_order_table",
            "asset_group": "wind_onshore",
            "fuel": "wind",
            "included": "true",
        },
    ]
    values = {
        "source_a_capacity_in_merit_order_table": {"future": 25.0},
        "source_a_operating_costs_in_merit_order_table": {"future": 4.0},
        "source_b_capacity_in_merit_order_table": {"future": 75.0},
        "source_b_operating_costs_in_merit_order_table": {"future": 8.0},
        "source_c_capacity_in_merit_order_table": {"future": 0.0},
    }

    assert aggregate_generation_costs(mappings, values) == {"solar_pv": 7.0}


def test_i_elgas_reformer_parameters_and_co2_cost_are_parsed(tmp_path) -> None:
    source = tmp_path / "technology.csv"
    source.write_text(
        "FUEL;TECH/PROCESS;EMISSIONS;Marginal Costs\n"
        "NGAS;ATR;5,68;4,94\n"
        "NGAS;SMR;56,8;5,4\n"
        "NGAS;SMR_CCS;31,24;9,73\n",
        encoding="utf-8",
    )

    parameters = load_i_elgas_technology_parameters(source)

    assert parameters["ATR"].emissions_kg_per_gj_input == 5.68
    assert parameters["SMR_CCS"].marginal_cost_eur_per_mwh_output == 9.73
    assert reformer_co2_cost_eur_per_mwh_output(5.68, 0.8, 100.0) == pytest.approx(
        2.556
    )


def test_static_generation_availability_is_capacity_weighted_and_profiled() -> None:
    mappings = [
        {"query_key": "a_capacity_in_merit_order_table", "asset_group": "nuclear", "fuel": "nuclear", "included": "true"},
        {"query_key": "b_capacity_in_merit_order_table", "asset_group": "nuclear", "fuel": "nuclear", "included": "true"},
    ]
    values = {
        "a_capacity_in_merit_order_table": {"future": 25.0},
        "a_availability_in_merit_order_table": {"future": 80.0},
        "b_capacity_in_merit_order_table": {"future": 75.0},
        "b_availability_in_merit_order_table": {"future": 90.0},
    }
    availability = aggregate_generation_availability(mappings, values)
    attachments, profiles = build_static_availability_profile_tables(
        [{"node": "E-A", "asset_group": "nuclear", "operating_mode": "must_run", "capacity_mw": 100.0}],
        2050,
        availability,
    )

    assert availability == {"nuclear": pytest.approx(0.875)}
    assert len(attachments) == 1
    assert len(profiles) == 8760
    assert {row["value"] for row in profiles} == {0.875}


def test_renewable_profiles_are_shared_and_solar_is_capacity_weighted(tmp_path) -> None:
    (tmp_path / "asset.csv").write_text(
        "asset,capacity\n"
        "NL00_Wind_Onshore,100\n"
        "NL00_Solar_Photovoltaic,25\n"
        "NL00_Solar_Rooftop,75\n",
        encoding="utf-8",
    )
    (tmp_path / "profiles-rep-periods.csv").write_text(
        "milestone_year,profile_name,rep_period,timestep,value\n"
        "2050,NL00_Wind_Onshore,1,1,0.4\n"
        "2050,NL00_Solar_Photovoltaic,1,1,0.8\n"
        "2050,NL00_Solar_Rooftop,1,1,0.4\n",
        encoding="utf-8",
    )
    rows = [
        {"node": "E-A", "asset_group": "wind_onshore", "operating_mode": "volatile", "capacity_mw": 10.0},
        {"node": "E-B", "asset_group": "wind_onshore", "operating_mode": "volatile", "capacity_mw": 20.0},
        {"node": "E-A", "asset_group": "solar_pv", "operating_mode": "volatile", "capacity_mw": 30.0},
    ]

    attachments, profiles = build_renewable_profile_tables(rows, 2050, tmp_path)

    assert len(attachments) == 3
    assert {row["profile_name"] for row in attachments} == {
        "SR2025_wind_onshore_2050",
        "SR2025_solar_pv_2050",
    }
    values = {row["profile_name"]: row["value"] for row in profiles}
    assert values["SR2025_wind_onshore_2050"] == 0.4
    assert values["SR2025_solar_pv_2050"] == pytest.approx(0.5)


def test_must_run_profile_and_minimum_energy_force_hourly_output() -> None:
    rows = [
        {
            "node": "E-A",
            "asset_group": "waste_steam",
            "operating_mode": "must_run",
            "capacity_mw": 20.0,
        },
        {
            "node": "E-B",
            "asset_group": "waste_steam",
            "operating_mode": "must_run",
            "capacity_mw": 80.0,
        },
    ]
    curve_rows = [
        {"energy_power_supercritical_waste_mix.output (MW)": 40.0},
        {"energy_power_supercritical_waste_mix.output (MW)": 75.0},
    ]

    attachments, profiles, minimum_energy = build_must_run_profile_tables(
        rows, 2050, curve_rows
    )

    assert len(attachments) == 2
    assert [row["value"] for row in profiles] == [0.4, 0.75]
    assert minimum_energy == {
        "E-A_waste_steam_must_run": pytest.approx(23.0),
        "E-B_waste_steam_must_run": pytest.approx(92.0),
    }


def test_coupled_assets_connect_nodal_electricity_to_country_carriers() -> None:
    electricity_rows = [
        {
            "node": "E-A",
            "asset_group": "hydrogen_ccgt",
            "fuel": "hydrogen",
            "operating_mode": "dispatchable",
            "capacity_mw": 20.0,
        },
        {
            "node": "E-B",
            "asset_group": "natural_gas_ccgt",
            "fuel": "natural_gas",
            "operating_mode": "dispatchable",
            "capacity_mw": 30.0,
        },
    ]
    hydrogen_rows = [
        {
            "asset_group": "electrolysis",
            "operating_mode": "mixed",
            "capacity_mw": 100.0,
        },
        {
            "asset_group": "atr_ccs",
            "operating_mode": "must_run",
            "capacity_mw": 40.0,
        },
    ]
    weights = {
        "electrolysis": {"E-A": 0.25, "E-B": 0.75},
        "atr_ccs": {"E-B": 1.0},
    }

    frames, reconciliation, _ = assemble_tulipa_tables(
        2030,
        electricity_rows,
        hydrogen_rows,
        weights,
        technology_parameters={
            group: TechnologyParameter(group, efficiency, cost, "test", "test")
            for group, efficiency, cost in (
                ("electrolysis", 0.7, 2.0),
                ("atr_ccs", 0.75, 3.0),
                ("hydrogen_ccgt", 0.6, 4.0),
                ("natural_gas_ccgt", 0.6, 5.0),
            )
        },
    )

    assets = {row["asset"]: row for row in frames["asset"]}
    assert {electricity_bus("E-A", 2030), electricity_bus("E-B", 2030)} <= assets.keys()
    assert {country_bus("hydrogen", 2030), country_bus("methane", 2030)} <= assets.keys()
    assert assets["E-A_electrolysis_mixed"]["capacity"] == pytest.approx(25.0)
    assert assets["E-B_electrolysis_mixed"]["capacity"] == pytest.approx(75.0)
    commissions = {row["asset"]: row for row in frames["asset-commission"]}
    assert commissions["E-A_electrolysis_mixed"]["conversion_efficiency"] == 0.7

    flows = {
        (row["from_asset"], row["to_asset"], row["carrier"])
        for row in frames["flow"]
    }
    assert (
        electricity_bus("E-A", 2030),
        "E-A_electrolysis_mixed",
        "electricity",
    ) in flows
    flow_commissions = {
        (row["from_asset"], row["to_asset"]): row
        for row in frames["flow-commission"]
    }
    assert flow_commissions[
        (electricity_bus("E-A", 2030), "E-A_electrolysis_mixed")
    ]["capacity_coefficient"] == 1.0
    assert flow_commissions[
        ("E-A_electrolysis_mixed", country_bus("hydrogen", 2030))
    ]["capacity_coefficient"] == 1.0
    assert (
        "E-A_electrolysis_mixed",
        country_bus("hydrogen", 2030),
        "hydrogen",
    ) in flows
    assert (
        country_bus("hydrogen", 2030),
        "E-A_hydrogen_ccgt_dispatchable",
        "hydrogen",
    ) in flows
    assert (
        country_bus("methane", 2030),
        "E-B_natural_gas_ccgt_dispatchable",
        "methane",
    ) in flows
    assert all(row["difference_mw"] == pytest.approx(0.0) for row in reconciliation)


def test_methane_supply_is_bounded_and_priced_upstream() -> None:
    frames, _, _ = assemble_tulipa_tables(
        2030,
        [],
        [],
        {},
        electricity_nodes=["E-A"],
        methane_rows=[
            {"supply_group": "fossil_methane", "summed_route_peak_mw": 20.0, "annual_supply_twh": 0.1},
            {"supply_group": "renewable_methane", "summed_route_peak_mw": 5.0, "annual_supply_twh": 0.02},
            {"supply_group": "network_gas_storage", "summed_route_peak_mw": 30.0, "annual_supply_twh": 0.2},
        ],
        fuel_prices_eur_per_mwh={"methane": 22.0, "renewable_methane": 64.0},
    )

    assets = {row["asset"]: row for row in frames["asset"]}
    assert assets["NL_fossil_methane_2030"]["capacity"] == 20.0
    assert assets["NL_renewable_methane_2030"]["capacity"] == 5.0
    assert "NL_network_gas_storage_2030" not in assets
    milestones = {row["asset"]: row for row in frames["asset-milestone"]}
    assert milestones["NL_fossil_methane_2030"]["max_energy_timeframe_partition"] == 100_000.0
    assert milestones["NL_renewable_methane_2030"]["max_energy_timeframe_partition"] == 20_000.0
    flow_milestones = {
        row["from_asset"]: row for row in frames["flow-milestone"]
    }
    assert flow_milestones["NL_fossil_methane_2030"]["commodity_price"] == 22.0
    assert flow_milestones["NL_renewable_methane_2030"]["commodity_price"] == 64.0


def test_demand_profiles_attach_to_nodal_and_national_consumers(tmp_path) -> None:
    electricity_path = tmp_path / "electricity.csv"
    electricity_path.write_text(
        "scenario_key,year,node,hour,demand_mw\n"
        + "".join(
            f"example,2030,E-A,{hour},{float(hour)}\n"
            for hour in range(1, 8761)
        ),
        encoding="utf-8",
    )
    grouped_path = tmp_path / "grouped.csv"
    grouped_path.write_text(
        "scenario_key,year,carrier,demand_group,hour,raw_mw,normalized_mw\n"
        + "".join(
            f"example,2030,{carrier},total,{hour},0,{value}\n"
            for carrier, value in (("hydrogen", 2.0), ("methane", 3.0))
            for hour in range(1, 8761)
        ),
        encoding="utf-8",
    )

    attachments, profiles = load_demand_profile_tables(
        "example", 2030, ["E-A"], electricity_path, grouped_path
    )

    assert {row["asset"] for row in attachments} == {
        electricity_bus("E-A", 2030),
        country_bus("hydrogen", 2030),
        country_bus("methane", 2030),
    }
    assert {row["profile_type"] for row in attachments} == {"demand"}
    assert len(profiles) == 3 * 8760
    methane_profile = f"SR2025_{country_bus('methane', 2030)}"
    assert {
        row["value"] for row in profiles if row["profile_name"] == methane_profile
    } == {3.0}


def test_electrolysis_efficiency_uses_hydrogen_output_capacity(tmp_path) -> None:
    workbook = Workbook()
    municipality = workbook.active
    municipality.title = "Municipality (ELEC, capacity)"
    industry = workbook.create_sheet("Industry (ELEC, capacity)")
    for worksheet in (municipality, industry):
        worksheet.append((None, "scenario", "scenario"))
        worksheet.append((None, 2050, 2050))
        worksheet.append((None, "Electricity", "Electricity"))
        worksheet.append((None, "Flexibility", "Flexibility"))
        worksheet.append((None, "Power_to_gas_onshore", "Power_to_gas_offshore"))
        worksheet.append(("code", None, None))
    municipality.append(("GM1", 13_000.0, 15_000.0))
    workbook_path = tmp_path / "scenario.xlsx"
    workbook.save(workbook_path)

    row = calculate_electrolysis_efficiency(
        "example",
        2050,
        [{"asset_group": "electrolysis", "capacity_mw": 21_212.121212}],
        workbook_path,
    )

    assert row is not None
    assert row.input_capacity_mw == 28_000.0
    assert row.output_capacity_mw == pytest.approx(21_212.121212)
    assert row.conversion_efficiency == pytest.approx(0.7575757576)


def test_generated_rows_follow_tyndp_table_columns() -> None:
    frames, _, _ = assemble_tulipa_tables(
        2030,
        [{
            "node": "E-A",
            "asset_group": "natural_gas_ccgt",
            "fuel": "natural_gas",
            "operating_mode": "dispatchable",
            "capacity_mw": 1.0,
        }],
        [],
        {},
        technology_parameters={
            "natural_gas_ccgt": TechnologyParameter(
                "natural_gas_ccgt", 0.6, 0.0, "test", "test"
            )
        },
    )

    for table, rows in frames.items():
        assert all(set(row) <= set(TABLE_COLUMNS[table]) for row in rows)

    assert all(
        row["use_binary_storage_method"] == "" for row in frames["asset"]
    )
    assert all(
        row["investment_limit"] == ""
        and row["investment_limit_storage_energy"] == ""
        for row in frames["asset-commission"]
    )
    assert all(
        row["initial_storage_level"] == "" and row["units_on_cost"] == ""
        for row in frames["asset-milestone"]
    )
    assert all(
        row["investment_limit"] == "" for row in frames["flow-commission"]
    )
    assert all(row["reactance"] == "" for row in frames["flow-milestone"])


def test_non_gas_electricity_generation_is_exported_as_producers() -> None:
    frames, _, _ = assemble_tulipa_tables(
        2050,
        [
            {
                "node": "E-A",
                "asset_group": "wind_onshore",
                "fuel": "wind",
                "operating_mode": "volatile",
                "capacity_mw": 120.0,
            },
            {
                "node": "E-B",
                "asset_group": "coal_steam_ccs",
                "fuel": "coal",
                "operating_mode": "dispatchable",
                "capacity_mw": 80.0,
            },
        ],
        [],
        {},
        generation_costs_eur_per_mwh={
            "wind_onshore": 1.5,
            "coal_steam_ccs": 72.0,
        },
    )

    assets = {row["asset"]: row for row in frames["asset"]}
    assert assets["E-A_wind_onshore_volatile"]["type"] == "producer"
    assert assets["E-A_wind_onshore_volatile"]["capacity"] == 120.0
    assert assets["E-B_coal_steam_ccs_dispatchable"]["type"] == "producer"
    assert assets["E-B_coal_steam_ccs_dispatchable"]["capacity"] == 80.0
    costs = {
        row["from_asset"]: row["operational_cost"]
        for row in frames["flow-milestone"]
    }
    assert costs["E-A_wind_onshore_volatile"] == 1.5
    assert costs["E-B_coal_steam_ccs_dispatchable"] == 72.0
    assert not any(
        row["to_asset"] in {
            "E-A_wind_onshore_volatile",
            "E-B_coal_steam_ccs_dispatchable",
        }
        for row in frames["flow"]
    )


def test_transport_link_uses_tulipa_bidirectional_capacity_structure() -> None:
    frames, _, _ = assemble_tulipa_tables(
        2025,
        [{
            "node": "E-A",
            "asset_group": "natural_gas_ccgt",
            "fuel": "natural_gas",
            "operating_mode": "dispatchable",
            "capacity_mw": 1.0,
        }],
        [],
        {},
        {"E-A", "E-B"},
        [{
            "from_endpoint": "E-A",
            "to_endpoint": "GER",
            "forward_capacity_mw": 1000.0,
            "reverse_capacity_mw": 750.0,
        }],
        technology_parameters={
            "natural_gas_ccgt": TechnologyParameter(
                "natural_gas_ccgt", 0.6, 0.0, "test", "test"
            )
        },
    )

    assets = {row["asset"] for row in frames["asset"]}
    assert "GER_E_Demand_2025" in assets
    transport = [row for row in frames["flow"] if row["is_transport"]]
    assert transport == [{
        "capacity": 1000.0,
        "carrier": "electricity",
        "discount_rate": 0.0,
        "economic_lifetime": 1,
        "from_asset": "E-A_E_Demand_2025",
        "investment_integer": False,
        "is_transport": True,
        "technical_lifetime": 1,
        "to_asset": "GER_E_Demand_2025",
    }]
    assert frames["flow-both"] == [{
        "commission_year": 2025,
        "decommissionable": False,
        "from_asset": "E-A_E_Demand_2025",
        "initial_export_units": 1.0,
        "initial_import_units": 0.75,
        "milestone_year": 2025,
        "to_asset": "GER_E_Demand_2025",
    }]


def test_2025_transport_uses_2027_and_filters_non_nl_electricity(tmp_path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Inputs I-ELGAS DE"
    worksheet.append([
        "CountryA", "CountryB", "Link", "2027", "2030", "2035", "2040",
        "2050", "InterfaceOrd",
    ])
    worksheet.append(["E-A", "E-B", 1, 100.0, 200.0, 0.0, 0.0, 0.0, 1])
    worksheet.append(["E-B", "E-A", 1, 80.0, 180.0, 0.0, 0.0, 0.0, 2])
    worksheet.append(["E-A", "GER", 1, 70.0, 170.0, 0.0, 0.0, 0.0, 3])
    worksheet.append(["GER", "E-A", 1, 60.0, 160.0, 0.0, 0.0, 0.0, 4])
    worksheet.append(["E-A", "H-A", 1, 999.0, 999.0, 0.0, 0.0, 0.0, 5])
    worksheet.append(["BEL", "GER", 1, 999.0, 999.0, 0.0, 0.0, 0.0, 6])
    path = tmp_path / "trading.xlsx"
    workbook.save(path)

    links = load_electricity_transport_links(path, 2025)

    assert links == [
        {
            "from_endpoint": "E-A",
            "to_endpoint": "E-B",
            "forward_capacity_mw": 100.0,
            "reverse_capacity_mw": 80.0,
            "source_year": 2027,
            "link_scope": "internal_nl",
        },
        {
            "from_endpoint": "E-A",
            "to_endpoint": "GER",
            "forward_capacity_mw": 70.0,
            "reverse_capacity_mw": 60.0,
            "source_year": 2027,
            "link_scope": "nl_cross_border",
        },
    ]


def test_storage_assets_preserve_power_volume_efficiency_and_carrier() -> None:
    storage_rows = [
        StorageTechnologyRow(
            "battery_system", "electricity", 120.0, 120.0, 960.0,
            1.0, 0.85, False,
        ),
        StorageTechnologyRow(
            "ides_storage", "electricity", 60.0, 60.0, 1080.0,
            1.0, 0.70, False,
        ),
        StorageTechnologyRow(
            "mdes_storage", "electricity", 32.0, 32.0, 3200.0,
            1.0, 0.80, False,
        ),
        StorageTechnologyRow(
            "hydrogen_storage_salt_cavern", "hydrogen", 12.0, 12.0, 3000.0,
            1.0, 1.0, True,
        ),
        StorageTechnologyRow(
            "hydrogen_storage_depleted_gas_field", "hydrogen", 48.0, 48.0,
            12000.0, 1.0, 1.0, True,
        ),
    ]

    frames, _, reconciliation = assemble_tulipa_tables(
        2050,
        [],
        [],
        {},
        {"E-A", "E-B"},
        storage_rows=storage_rows,
        electricity_storage_weights={
            group: {"E-A": 0.25, "E-B": 0.75}
            for group in ("battery_system", "ides_storage", "mdes_storage")
        },
    )

    assets = {row["asset"]: row for row in frames["asset"]}
    commissions = {row["asset"]: row for row in frames["asset-commission"]}
    assert assets["E-A_battery_system_2050"]["capacity"] == 30.0
    assert assets["E-A_battery_system_2050"]["capacity_storage_energy"] == 240.0
    assert assets["E-A_battery_system_2050"]["energy_to_power_ratio"] == 8.0
    assert commissions["E-A_battery_system_2050"]["storage_discharging_efficiency"] == 0.85
    assert assets["NL_hydrogen_storage_salt_cavern_2050"]["is_seasonal"] is True
    assert assets["NL_hydrogen_storage_salt_cavern_2050"]["capacity"] == 12.0
    assert assets["NL_hydrogen_storage_salt_cavern_2050"]["capacity_storage_energy"] == 3000.0
    flows = {(row["from_asset"], row["to_asset"], row["carrier"]) for row in frames["flow"]}
    assert (electricity_bus("E-A", 2050), "E-A_battery_system_2050", "electricity") in flows
    assert ("NL_hydrogen_storage_salt_cavern_2050", country_bus("hydrogen", 2050), "hydrogen") in flows
    storage_assets = {
        "E-A_battery_system_2050",
        "E-A_ides_storage_2050",
        "E-A_mdes_storage_2050",
        "NL_hydrogen_storage_salt_cavern_2050",
        "NL_hydrogen_storage_depleted_gas_field_2050",
    }
    assert all(
        row["reactance"] == ""
        for row in frames["flow-milestone"]
        if row["from_asset"] in storage_assets or row["to_asset"] in storage_assets
    )
    assert len(reconciliation) == 5
    assert all(row["input_difference_mw"] == pytest.approx(0.0) for row in reconciliation)
    assert all(row["output_difference_mw"] == pytest.approx(0.0) for row in reconciliation)
    assert all(row["storage_volume_difference_mwh"] == pytest.approx(0.0) for row in reconciliation)


def test_imported_hydrogen_carriers_are_separate_output_based_producers() -> None:
    hydrogen_rows = [
        {"asset_group": "ammonia_reformer", "operating_mode": "dispatchable", "capacity_mw": 50.0},
        {"asset_group": "ammonia_reformer", "operating_mode": "must_run", "capacity_mw": 150.0},
        {"asset_group": "lohc_reformer", "operating_mode": "unspecified", "capacity_mw": 100.0},
        {"asset_group": "liquid_hydrogen_regasifier", "operating_mode": "unspecified", "capacity_mw": 200.0},
    ]

    frames, reconciliation, _ = assemble_tulipa_tables(
        2030,
        [],
        hydrogen_rows,
        {},
        {"E-A"},
        hydrogen_carrier_import_costs_eur_per_mwh={
            "ammonia_reformer": 240.0,
            "lohc_reformer": 260.0,
            "liquid_hydrogen_regasifier": 203.0,
        },
        hydrogen_carrier_production_twh={
            ("ammonia_reformer", "must_run"): 1.25,
        },
    )

    assets = {row["asset"]: row for row in frames["asset"]}
    expected = {
        "NL_ammonia_reformer_dispatchable_2030": (50.0, 240.0),
        "NL_ammonia_reformer_must_run_2030": (150.0, 240.0),
        "NL_lohc_reformer_unspecified_2030": (100.0, 260.0),
        "NL_liquid_hydrogen_regasifier_unspecified_2030": (200.0, 203.0),
    }
    flow_costs = {
        row["from_asset"]: row["operational_cost"]
        for row in frames["flow-milestone"]
    }
    for asset, (capacity, cost) in expected.items():
        assert assets[asset]["type"] == "producer"
        assert assets[asset]["capacity"] == capacity
        assert flow_costs[asset] == cost
    assert sum(float(row["national_capacity_mw"]) for row in reconciliation) == 500.0
    assert all(row["difference_mw"] == 0.0 for row in reconciliation)
    milestones = {row["asset"]: row for row in frames["asset-milestone"]}
    assert milestones["NL_ammonia_reformer_must_run_2030"][
        "min_energy_timeframe_partition"
    ] == 1_250_000.0
    assert milestones["NL_ammonia_reformer_dispatchable_2030"][
        "min_energy_timeframe_partition"
    ] == ""


def test_direct_hydrogen_imports_preserve_capacity_cost_and_baseload_energy() -> None:
    hydrogen_rows = [
        {"asset_group": "hydrogen_import_backup", "operating_mode": "backup", "capacity_mw": 250.0},
        {"asset_group": "hydrogen_import_baseload", "operating_mode": "baseload", "capacity_mw": 100.0},
    ]

    frames, reconciliation, _ = assemble_tulipa_tables(
        2050,
        [],
        hydrogen_rows,
        {},
        {"E-A"},
        hydrogen_import_costs_eur_per_mwh={
            "hydrogen_import_backup": 126.0,
            "hydrogen_import_baseload": 106.0,
        },
    )

    assets = {row["asset"]: row for row in frames["asset"]}
    assert assets["NL_hydrogen_import_backup_2050"]["capacity"] == 250.0
    assert assets["NL_hydrogen_import_baseload_2050"]["capacity"] == 100.0
    costs = {
        row["from_asset"]: row["operational_cost"]
        for row in frames["flow-milestone"]
    }
    assert costs["NL_hydrogen_import_backup_2050"] == 126.0
    assert costs["NL_hydrogen_import_baseload_2050"] == 106.0
    milestones = {row["asset"]: row for row in frames["asset-milestone"]}
    assert milestones["NL_hydrogen_import_backup_2050"]["min_energy_timeframe_partition"] == ""
    assert milestones["NL_hydrogen_import_baseload_2050"]["min_energy_timeframe_partition"] == 876000.0
    assert {
        row["asset_group"] for row in reconciliation
    } >= {"hydrogen_import_backup", "hydrogen_import_baseload"}